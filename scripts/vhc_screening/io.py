"""Reading and validation of InvestingPro exports."""

import re
from collections.abc import Callable, Sequence
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

import openpyxl
from defusedxml.common import DefusedXmlException
from openpyxl.utils.exceptions import InvalidFileException

from . import i18n
from .constants import (
    CORE_COLUMNS,
    ESSENTIAL_COLUMNS,
    NA_VALUES,
    NAME,
    PERCENTAGE_COLUMNS,
    TICKER,
)

Row = Sequence[Any]
HeaderMap = dict[str, int]
Accessor = Callable[[Row, str], Any]

# Defensive limits: a normal InvestingPro export holds up to 1,000 rows.
MAX_XLSX_FILE_BYTES = 5 * 1024 * 1024  # 5 MiB for the attached XLSX file.
MAX_XLSX_ARCHIVE_ENTRIES = 2_048  # 2,048 internal files at most.
MAX_XLSX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024  # 200 MiB in total.
MAX_XLSX_MEMBER_BYTES = 100 * 1024 * 1024  # 100 MiB per internal file.
MAX_WORKSHEET_ROWS = 2_000  # 2,000 rows: twice the expected export.
MAX_WORKSHEET_COLUMNS = 256  # 256 columns per sheet.
SECTOR_LIST_MAX_AGE_DAYS = 180  # Six months between CSV reviews.

# Extracts the date from comments such as: "corte 2026-07-13".
SECTOR_CUTOFF_PATTERN = re.compile(r"\bcorte\s+(\d{4}-\d{2}-\d{2})\b", re.IGNORECASE)


def normalize(value: Any) -> str:
    """Normalizes whitespace and line breaks in a header."""
    return " ".join(str(value).split()) if value is not None else ""


def _find_header_row(matrix: list[list[Any]], language) -> int:
    for index, row in enumerate(matrix):
        values = {normalize(value) for value in row}
        if NAME in values and TICKER in values:
            return index
    raise SystemExit(i18n.text("io_no_header", language))


def _validate_export_path(path: str | Path, language) -> Path:
    export_path = Path(path)
    if export_path.suffix.lower() != ".xlsx":
        raise SystemExit(i18n.text("io_not_xlsx", language))
    try:
        if not export_path.is_file():
            raise SystemExit(i18n.text("io_not_found", language, path=export_path))
        size = export_path.stat().st_size
    except OSError as error:
        raise SystemExit(i18n.text("io_stat_failed", language, error=error)) from error
    if size == 0:
        raise SystemExit(i18n.text("io_empty", language))
    if size > MAX_XLSX_FILE_BYTES:
        limit_mb = MAX_XLSX_FILE_BYTES // (1024 * 1024)
        raise SystemExit(i18n.text("io_too_large", language, limit=limit_mb))
    return export_path


def _inspect_xlsx_archive(path: Path, language) -> None:
    """Validates basic ZIP container limits before parsing its XML."""
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_ARCHIVE_ENTRIES:
                raise SystemExit(i18n.text("io_too_many_members", language))

            names = {member.filename for member in members}
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise SystemExit(i18n.text("io_bad_structure", language))

            if any(member.file_size > MAX_XLSX_MEMBER_BYTES for member in members):
                raise SystemExit(i18n.text("io_member_too_large", language))
            total_size = sum(member.file_size for member in members)
            if total_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise SystemExit(i18n.text("io_uncompressed_too_large", language))
    except BadZipFile as error:
        raise SystemExit(i18n.text("io_corrupt", language)) from error


def _read_sheet_rows(sheet, language) -> list[list[Any]]:
    if sheet.max_row is not None and sheet.max_row > MAX_WORKSHEET_ROWS:
        raise SystemExit(
            i18n.text("io_too_many_rows", language, limit=f"{MAX_WORKSHEET_ROWS:,}")
        )
    if sheet.max_column is not None and sheet.max_column > MAX_WORKSHEET_COLUMNS:
        raise SystemExit(
            i18n.text("io_too_many_columns", language, limit=MAX_WORKSHEET_COLUMNS)
        )

    matrix = []
    for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if row_index > MAX_WORKSHEET_ROWS:
            raise SystemExit(
                i18n.text("io_too_many_rows", language, limit=f"{MAX_WORKSHEET_ROWS:,}")
            )
        values = list(row)
        if len(values) > MAX_WORKSHEET_COLUMNS:
            raise SystemExit(
                i18n.text("io_too_many_columns", language, limit=MAX_WORKSHEET_COLUMNS)
            )
        matrix.append(values)
    return matrix


def read_export(
    path: str | Path, language: str = i18n.SPANISH
) -> tuple[list[list[Any]], HeaderMap, list[str]]:
    """Reads the export by column name, regardless of column order."""
    export_path = _validate_export_path(path, language)
    _inspect_xlsx_archive(export_path, language)
    try:
        workbook = openpyxl.load_workbook(
            export_path,
            data_only=True,
            read_only=True,
            keep_links=False,
        )
        try:
            matrix = _read_sheet_rows(workbook.active, language)
        finally:
            workbook.close()
    except (
        BadZipFile,
        DefusedXmlException,
        EOFError,
        IndexError,
        InvalidFileException,
        KeyError,
        OSError,
        TypeError,
        ValueError,
    ) as error:
        raise SystemExit(i18n.text("io_unreadable", language, error=error)) from error

    header_row = _find_header_row(matrix, language)
    header_map = {
        normalize(value): index
        for index, value in enumerate(matrix[header_row])
        if normalize(value)
    }
    missing_essential = [
        column for column in ESSENTIAL_COLUMNS if column not in header_map
    ]
    if missing_essential:
        raise SystemExit(
            i18n.text("io_missing_essential", language, columns=missing_essential)
        )

    missing_core = [column for column in CORE_COLUMNS if column not in header_map]
    name_index = header_map[NAME]
    rows = [
        row for row in matrix[header_row + 1 :] if row and normalize(row[name_index])
    ]
    return rows, header_map, missing_core


def load_sector_exclusions(
    path: str | Path, language: str = i18n.SPANISH
) -> dict[str, str]:
    """Loads the mandatory list of sectors this methodology does not evaluate."""
    csv_path = Path(path)
    if not csv_path.exists():
        raise SystemExit(i18n.text("io_sector_missing", language, path=csv_path))

    sectors: dict[str, str] = {}
    try:
        with open(csv_path, encoding="utf-8") as source:
            for raw_line in source:
                line = raw_line.strip()
                if not line or line.startswith(("#", "ticker,")):
                    continue
                parts = line.split(",")
                if len(parts) >= 3:
                    sectors[parts[0].strip()] = parts[-1].strip()
    except OSError as error:
        raise SystemExit(
            i18n.text("io_sector_unreadable", language, path=csv_path, error=error)
        ) from error

    if not sectors:
        raise SystemExit(i18n.text("io_sector_malformed", language, path=csv_path))
    return sectors


def sector_list_age_warning(
    path: str | Path,
    today: date | None = None,
    max_age_days: int = SECTOR_LIST_MAX_AGE_DAYS,
    language: str = i18n.SPANISH,
) -> str | None:
    """Warns when the sector CSV cut-off date needs review."""
    csv_path = Path(path)
    try:
        first_lines = "".join(csv_path.read_text(encoding="utf-8").splitlines(True)[:5])
    except OSError:
        return None

    match = SECTOR_CUTOFF_PATTERN.search(first_lines)
    if not match:
        return i18n.text("sector_no_cutoff", language)
    try:
        cutoff = date.fromisoformat(match.group(1))
    except ValueError:
        return i18n.text("sector_bad_cutoff", language)
    current_date = today or datetime.now(UTC).astimezone().date()
    age_days = (current_date - cutoff).days
    if age_days <= max_age_days:
        return None
    return i18n.text("sector_stale", language, days=age_days, cutoff=cutoff.isoformat())


def verify_percentage_scale(rows: Sequence[Row], header_map: HeaderMap) -> list[str]:
    """Detects columns that appear to have switched from decimal to percentage."""
    suspicious = []
    for column in PERCENTAGE_COLUMNS:
        index = header_map.get(column)
        if index is None:
            continue
        values = sorted(
            abs(value)
            for value in (row[index] if index < len(row) else None for row in rows)
            if isinstance(value, (int, float))
        )
        if len(values) >= 20 and values[len(values) // 2] >= 5:
            suspicious.append(column)
    return suspicious


def make_accessors(header_map: HeaderMap) -> tuple[Accessor, Accessor, Accessor]:
    """Creates raw, numeric and percentage accessors for a schema."""

    def raw(row: Row, header: str) -> Any:
        index = header_map.get(header)
        if index is None or index >= len(row):
            return None
        value = row[index]
        if value is None:
            return None
        if isinstance(value, str) and value.strip().upper() in NA_VALUES:
            return None
        return value

    def number(row: Row, header: str) -> float | None:
        value = raw(row, header)
        return float(value) if isinstance(value, (int, float)) else None

    def percentage(row: Row, header: str) -> float | None:
        # InvestingPro delivers these percentages as decimals: 0.478 = 47.8 %.
        value = number(row, header)
        return None if value is None else value * 100

    return raw, number, percentage
