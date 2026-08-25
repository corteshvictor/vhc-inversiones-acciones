"""Enriched Excel workbook generation."""

from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import WorksheetProperties
from openpyxl.worksheet.worksheet import Worksheet

from . import i18n
from .constants import (
    ALERT_PREFIX,
    ALERT_ROIC,
    ALERT_VALUE_TRAP,
    ALTMAN,
    BENEISH,
    BETA,
    BUCKET_DEEP_DIVE,
    BUCKET_DISCARDED,
    BUCKET_EMOJI,
    BUCKET_NEUTRAL,
    BUCKET_WATCHLIST,
    BUYBACK_YIELD,
    CORE_COLUMNS,
    CURRENT_PRICE,
    DEBT_CAPITAL,
    DEFAULT_ASSETS,
    EPS_GROWTH,
    EV_EBIT,
    FAIR_VALUE,
    FAIR_VALUE_LABEL,
    FCF_NET_INCOME,
    FCF_YIELD,
    GROSS_MARGIN,
    HEALTH_LABEL,
    MARKET_CAP,
    NAME,
    PE_RATIO,
    PEG_RATIO,
    PIOTROSKI,
    PRICE_CHEAP,
    PRICE_EXPENSIVE,
    PRICE_MIXED,
    PRICE_VERY_CHEAP,
    QUALITY_EXCELLENT,
    QUALITY_GOOD,
    QUALITY_WEAK,
    REVENUE_GROWTH,
    ROE,
    ROIC,
    ROIC_5Y,
    SCREENING_VERSION,
    TICKER,
    VERDICT_MEDIUM,
    VERDICT_NO,
    VERDICT_YES,
)
from .io import HeaderMap, Row, make_accessors

RESULT = "RESULTADO"
COMPANY = "EMPRESA"
QUALITY = "CALIDAD"
PRICE = "PRECIO"
HEALTH = "SALUD"
BUCKET = "Balde"
SCORE = "Puntaje (0-10)"
GREENBLATT = "Greenblatt"
MSCI = "MSCI"
AQR = "AQR"
QUALITY_VERDICT = "Veredicto Calidad"
QUALITY_SIGNALS = "Señales de calidad (0-3)"
UPSIDE = "Upside vs Fair Value"
PRICE_VERDICT = "Veredicto Precio"
PRICE_SIGNALS = "Señales de precio (0-4)"
HEALTH_VERDICT = "Veredicto Salud financiera"
HEALTH_SIGNALS = "Señales de salud (0-3)"
ALERTS = "Alertas"

FONT_NAME = "Arial"
SOLID = "solid"
WHITE = "FFFFFF"
DARK_TEXT = "262626"
CENTER = "center"
METHOD = "Método"
FUND_LABEL = "Fondo/ETF"
OMITTED_LABEL = "Omitida"
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")

LAYOUT = [
    (BUCKET, RESULT, True),
    (SCORE, RESULT, True),
    (NAME, COMPANY, False),
    (TICKER, COMPANY, False),
    (CURRENT_PRICE, COMPANY, False),
    (MARKET_CAP, COMPANY, False),
    (ROIC, QUALITY, False),
    (ROIC_5Y, QUALITY, False),
    (ROE, QUALITY, False),
    (GROSS_MARGIN, QUALITY, False),
    (EPS_GROWTH, QUALITY, False),
    (REVENUE_GROWTH, QUALITY, False),
    (FCF_NET_INCOME, QUALITY, False),
    (BUYBACK_YIELD, QUALITY, False),
    (DEBT_CAPITAL, QUALITY, False),
    (BETA, QUALITY, False),
    (GREENBLATT, QUALITY, True),
    (MSCI, QUALITY, True),
    (AQR, QUALITY, True),
    (QUALITY_VERDICT, QUALITY, True),
    (QUALITY_SIGNALS, QUALITY, True),
    (PE_RATIO, PRICE, False),
    (PEG_RATIO, PRICE, False),
    (EV_EBIT, PRICE, False),
    (FCF_YIELD, PRICE, False),
    (FAIR_VALUE, PRICE, False),
    (FAIR_VALUE_LABEL, PRICE, False),
    (UPSIDE, PRICE, True),
    (PRICE_VERDICT, PRICE, True),
    (PRICE_SIGNALS, PRICE, True),
    (PIOTROSKI, HEALTH, False),
    (ALTMAN, HEALTH, False),
    (BENEISH, HEALTH, False),
    (HEALTH_LABEL, HEALTH, False),
    (HEALTH_VERDICT, HEALTH, True),
    (HEALTH_SIGNALS, HEALTH, True),
    (ALERTS, RESULT, True),
]

BAND_COLORS = {
    RESULT: "404040",
    COMPANY: "595959",
    QUALITY: "375623",
    PRICE: "7F6000",
    HEALTH: "1F4E79",
}
DARK_COLORS = {
    RESULT: "404040",
    COMPANY: "595959",
    QUALITY: "538135",
    PRICE: "BF8F00",
    HEALTH: "2E75B6",
}
LIGHT_COLORS = {
    RESULT: "D9D9D9",
    COMPANY: "D9D9D9",
    QUALITY: "C6E0B4",
    PRICE: "FFE699",
    HEALTH: "BDD7EE",
}
BAND_LABELS = {
    RESULT: RESULT,
    COMPANY: COMPANY,
    QUALITY: "CALIDAD — ¿es un gran negocio?",
    PRICE: "PRECIO — ¿está barata?",
    HEALTH: "SALUD FINANCIERA — ¿sólida y limpia?",
}
PERCENT_COLUMNS = {
    FCF_YIELD,
    ROIC,
    ROIC_5Y,
    ROE,
    GROSS_MARGIN,
    EPS_GROWTH,
    REVENUE_GROWTH,
    FCF_NET_INCOME,
    BUYBACK_YIELD,
    DEBT_CAPITAL,
    UPSIDE,
}
TWO_DECIMAL_COLUMNS = {
    CURRENT_PRICE,
    FAIR_VALUE,
    PE_RATIO,
    PEG_RATIO,
    EV_EBIT,
    ALTMAN,
    BENEISH,
    BETA,
}
CENTER_COLUMNS = {
    GREENBLATT,
    MSCI,
    AQR,
    QUALITY_SIGNALS,
    PRICE_SIGNALS,
    HEALTH_SIGNALS,
    SCORE,
}

HEADER_COMMENTS = {
    SCORE: "xl_score",
    QUALITY_VERDICT: "xl_quality_verdict",
    PRICE_VERDICT: "xl_price_verdict",
    HEALTH_VERDICT: "xl_health_verdict",
    PIOTROSKI: "xl_piotroski",
    ALTMAN: "xl_altman",
    BENEISH: "xl_beneish",
    ROIC: "xl_roic",
    ROIC_5Y: "xl_roic_5y",
    ROE: "xl_roe",
    GROSS_MARGIN: "xl_gross_margin",
    EV_EBIT: "xl_ev_ebit",
    FCF_YIELD: "xl_fcf_yield",
    FAIR_VALUE: "xl_fair_value",
    DEBT_CAPITAL: "xl_debt_capital",
}


@dataclass(frozen=True)
class CellStyles:
    base_font: Font
    red_fill: PatternFill
    red_font: Font
    amber_fill: PatternFill
    amber_font: Font
    green_fill: PatternFill
    green_font: Font
    strong_green_fill: PatternFill
    strong_green_font: Font
    bucket_fills: dict[str, PatternFill]


def _fill(color: str) -> PatternFill:
    return PatternFill(SOLID, start_color=color)


def _styles() -> CellStyles:
    return CellStyles(
        base_font=Font(name=FONT_NAME, size=10),
        red_fill=_fill("FFC7CE"),
        red_font=Font(name=FONT_NAME, size=10, color="9C0006"),
        amber_fill=_fill("FFEB9C"),
        amber_font=Font(name=FONT_NAME, size=10, color="9C6500"),
        green_fill=_fill("C6EFCE"),
        green_font=Font(name=FONT_NAME, size=10, color="006100"),
        strong_green_fill=_fill("A9D08E"),
        strong_green_font=Font(name=FONT_NAME, size=10, color="004B00", bold=True),
        bucket_fills={
            BUCKET_DEEP_DIVE: _fill("A9D08E"),
            BUCKET_WATCHLIST: _fill("E2EFDA"),
            BUCKET_NEUTRAL: _fill("FFF2CC"),
            BUCKET_DISCARDED: _fill("FFC7CE"),
            FUND_LABEL: _fill("E0E0E0"),
            OMITTED_LABEL: _fill("D6DEE8"),
        },
    )


def _write_band_headers(sheet, offset: int, language) -> None:
    start = 0
    while start < len(LAYOUT):
        group = LAYOUT[start][1]
        end = start
        while end < len(LAYOUT) and LAYOUT[end][1] == group:
            end += 1
        sheet.merge_cells(
            f"{get_column_letter(start + 1)}{1 + offset}:"
            f"{get_column_letter(end)}{1 + offset}"
        )
        cell = sheet.cell(
            1 + offset, start + 1, i18n.term(BAND_LABELS[group], language)
        )
        cell.fill = _fill(BAND_COLORS[group])
        cell.font = Font(name=FONT_NAME, bold=True, color=WHITE, size=10)
        cell.alignment = Alignment(horizontal=CENTER, vertical=CENTER)
        start = end


def _display_header(header: str, language) -> str:
    renamed = header.replace(MARKET_CAP, "Market Cap (US$ B)").replace(
        UPSIDE, "Price vs Fair Value"
    )
    return i18n.term(renamed, language)


def _write_column_headers(sheet, offset: int, language) -> dict[str, int]:
    names = {}
    for index, (header, group, calculated) in enumerate(LAYOUT, 1):
        names[header] = index
        cell = sheet.cell(2 + offset, index, _display_header(header, language))
        cell.fill = _fill(DARK_COLORS[group] if calculated else LIGHT_COLORS[group])
        cell.font = Font(
            name=FONT_NAME,
            bold=True,
            color=WHITE if calculated else DARK_TEXT,
            size=10,
        )
        cell.alignment = Alignment(horizontal=CENTER, vertical=CENTER, wrap_text=True)
    return names


def _add_header_comments(sheet, names: dict[str, int], offset: int, language) -> None:
    author = i18n.term(METHOD, language)
    for header, key in HEADER_COMMENTS.items():
        if header in names:
            comment = Comment(i18n.text(key, language), author)
            sheet.cell(2 + offset, names[header]).comment = comment


def _has_declining_roic(evaluation: dict) -> bool:
    if evaluation["fund"]:
        return False
    metrics = evaluation["metrics"]
    return (
        metrics["roic"] is not None
        and metrics["roic5"] is not None
        and metrics["roic"] < metrics["roic5"]
    )


def _raw_value(row: Row, header: str, raw) -> object:
    value = raw(row, header)
    if header == MARKET_CAP and isinstance(value, (int, float)):
        return value / 1e9
    return value


def _excel_safe_value(value):
    """Prevents text coming from the export from being read as a formula."""
    if not isinstance(value, str):
        return value
    candidate = value.lstrip(" ")
    return f"'{value}" if candidate.startswith(FORMULA_PREFIXES) else value


def _fund_value(
    evaluation: dict, row: Row, header: str, calculated: bool, raw, language
):
    sector = evaluation.get("excluded_sector")
    if header == BUCKET:
        if not sector:
            return i18n.term(FUND_LABEL, language)
        return i18n.text("xl_omitted_cell", language)
    if header == ALERTS:
        if not sector:
            return i18n.text("xl_fund_cell", language)
        return i18n.text(
            "xl_omitted_comment", language, sector=i18n.term(sector, language)
        )
    return None if calculated else _raw_value(row, header, raw)


CALCULATED_FIELDS = {
    SCORE: lambda item: item["score"],
    QUALITY_VERDICT: lambda item: item["quality"],
    QUALITY_SIGNALS: lambda item: item["quality_count"],
    UPSIDE: lambda item: item["upside"],
    PRICE_VERDICT: lambda item: item["price"],
    PRICE_SIGNALS: lambda item: item["price_count"],
    HEALTH_VERDICT: lambda item: item["health"],
    HEALTH_SIGNALS: lambda item: item["health_count"],
}


def _alerts_value(evaluation: dict, declining_roic: bool, language) -> str:
    value = evaluation["alerts"]
    actionable = evaluation["bucket"] in {BUCKET_DEEP_DIVE, BUCKET_WATCHLIST}
    if declining_roic and actionable:
        value = ALERT_ROIC if value == "—" else f"{value} · {ALERT_ROIC}"
    if value == "—":
        return value
    return " · ".join(
        f"{ALERT_PREFIX.get(alert, '')}{i18n.term(alert, language)}"
        for alert in value.split(" · ")
    )


def _company_value(
    evaluation: dict, row: Row, header: str, raw, declining_roic: bool, language
):
    if header == BUCKET:
        return BUCKET_EMOJI.get(evaluation["bucket"], evaluation["bucket"])
    if header in {GREENBLATT, MSCI, AQR}:
        return evaluation["lenses"][[GREENBLATT, MSCI, AQR].index(header)]
    if header == ALERTS:
        return _alerts_value(evaluation, declining_roic, language)
    getter = CALCULATED_FIELDS.get(header)
    return getter(evaluation) if getter else _raw_value(row, header, raw)


def _localized(value, language):
    """Renders a cell value in English, keeping any emoji prefix in place."""
    if language == i18n.SPANISH or not isinstance(value, str):
        return value
    direct = i18n.TERMS.get(value)
    if direct is not None:
        return direct
    prefix, _, rest = value.partition(" ")
    if rest in i18n.TERMS:
        return f"{prefix} {i18n.TERMS[rest]}"
    return value


def _cell_colors(
    evaluation: dict,
    header: str,
    value,
    declining_roic: bool,
    styles: CellStyles,
) -> tuple[PatternFill | None, Font]:
    if evaluation["fund"]:
        return None, styles.base_font
    if header in evaluation["red_flags"] or (
        header == UPSIDE and evaluation["upside_red"]
    ):
        return styles.red_fill, styles.red_font
    if header in evaluation["amber_flags"] or (header == ROIC and declining_roic):
        return styles.amber_fill, styles.amber_font

    status_colors = {
        VERDICT_YES: (styles.green_fill, styles.green_font),
        VERDICT_MEDIUM: (styles.amber_fill, styles.amber_font),
        VERDICT_NO: (styles.red_fill, styles.red_font),
        QUALITY_EXCELLENT: (styles.green_fill, styles.green_font),
        QUALITY_GOOD: (styles.amber_fill, styles.amber_font),
        QUALITY_WEAK: (styles.red_fill, styles.red_font),
        PRICE_VERY_CHEAP: (styles.strong_green_fill, styles.strong_green_font),
        PRICE_CHEAP: (styles.green_fill, styles.green_font),
        PRICE_EXPENSIVE: (styles.red_fill, styles.red_font),
        PRICE_MIXED: (styles.amber_fill, styles.amber_font),
    }
    if header in {
        GREENBLATT,
        MSCI,
        AQR,
        QUALITY_VERDICT,
        PRICE_VERDICT,
        HEALTH_VERDICT,
    }:
        return status_colors.get(value, (None, styles.base_font))
    if header == ALERTS and value != "—":
        return styles.amber_fill, styles.amber_font
    return None, styles.base_font


def _format_cell(cell, header: str, value) -> None:
    if header in PERCENT_COLUMNS and isinstance(value, (int, float)):
        cell.number_format = "0.0%"
    elif header == MARKET_CAP and isinstance(value, (int, float)):
        cell.number_format = "#,##0.0"
    elif header in TWO_DECIMAL_COLUMNS and isinstance(value, (int, float)):
        cell.number_format = "0.00"
    if header in CENTER_COLUMNS:
        cell.alignment = Alignment(horizontal=CENTER)


def _source_value(evaluation, source_row, header, calculated, raw, declining, language):
    """The value a cell carries, before any color or format is decided."""
    if evaluation["fund"]:
        return _fund_value(evaluation, source_row, header, calculated, raw, language)
    return _company_value(evaluation, source_row, header, raw, declining, language)


def _bucket_fill_key(evaluation: dict) -> str:
    """Which fill the first cell of the row takes."""
    if evaluation.get("excluded_sector"):
        return OMITTED_LABEL
    if evaluation["fund"]:
        return FUND_LABEL
    return evaluation["bucket"]


def _write_data_rows(
    sheet, rows, evaluations, order, raw, offset, styles, language
) -> int:
    row_index = 3 + offset
    for key in order:
        evaluation = evaluations[key]
        source_row = rows[key]
        declining_roic = _has_declining_roic(evaluation)
        for column_index, (header, _group, calculated) in enumerate(LAYOUT, 1):
            value = _source_value(
                evaluation,
                source_row,
                header,
                calculated,
                raw,
                declining_roic,
                language,
            )
            # Only what the method computes may be translated. A raw export
            # value that happened to read "Calidad" is InvestingPro's data and
            # must reach the sheet untouched.
            shown = _localized(value, language) if calculated else value
            cell = sheet.cell(row_index, column_index, _excel_safe_value(shown))
            fill, font = _cell_colors(evaluation, header, value, declining_roic, styles)
            cell.font = font
            if fill is not None:
                cell.fill = fill
            _format_cell(cell, header, value)
        sheet.cell(row_index, 1).fill = styles.bucket_fills[
            _bucket_fill_key(evaluation)
        ]
        row_index += 1
    return row_index


def _mark_missing_columns(sheet, missing, names, offset, last_row, language) -> None:
    if not missing:
        return
    sheet.merge_cells(f"A1:{get_column_letter(len(LAYOUT))}1")
    banner = sheet.cell(
        1,
        1,
        i18n.text(
            "xl_missing_banner",
            language,
            count=len(missing),
            columns=" · ".join(missing),
        ),
    )
    banner.fill = _fill("C00000")
    banner.font = Font(name=FONT_NAME, bold=True, color=WHITE, size=11)
    banner.alignment = Alignment(horizontal="left", vertical=CENTER, wrap_text=True)
    sheet.row_dimensions[1].height = 32

    missing_fill = _fill("FFC7CE")
    missing_font = Font(name=FONT_NAME, size=10, color="9C0006", italic=True)
    for header in missing:
        column_index = names[header]
        header_cell = sheet.cell(2 + offset, column_index)
        header_cell.value = f"{header}  ⚠️"
        header_cell.fill = _fill("C00000")
        header_cell.font = Font(name=FONT_NAME, bold=True, color=WHITE, size=10)
        header_cell.comment = Comment(
            i18n.text("xl_missing_comment", language, column=header),
            i18n.term(METHOD, language),
        )
        placeholder = "missing" if language == i18n.ENGLISH else "falta"
        for row_index in range(3 + offset, last_row):
            cell = sheet.cell(row_index, column_index, placeholder)
            cell.fill = missing_fill
            cell.font = missing_font
            cell.alignment = Alignment(horizontal=CENTER, vertical=CENTER)


def _configure_screening_sheet(sheet, offset, names, last_row) -> None:
    sheet.freeze_panes = f"E{3 + offset}"
    sheet.auto_filter.ref = (
        f"A{2 + offset}:{get_column_letter(len(LAYOUT))}{last_row - 1}"
    )
    upside_letter = get_column_letter(names[UPSIDE])
    data_range = f"{upside_letter}{3 + offset}:{upside_letter}{last_row - 1}"
    sheet.conditional_formatting.add(
        data_range,
        DataBarRule(
            start_type="min",
            end_type="max",
            color="63C384",
            showValue=True,
        ),
    )
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 30
    for column_index in range(4, len(LAYOUT) + 1):
        sheet.column_dimensions[get_column_letter(column_index)].width = 11
    alerts_letter = get_column_letter(len(LAYOUT))
    sheet.column_dimensions[alerts_letter].width = 46
    for row_index in range(3 + offset, last_row):
        sheet[f"{alerts_letter}{row_index}"].alignment = Alignment(
            wrap_text=True,
            vertical=CENTER,
        )
    if not offset:
        sheet.row_dimensions[1].height = 18
    sheet.row_dimensions[2 + offset].height = 30


def _summarize(evaluations):
    companies = [key for key, value in evaluations.items() if not value["fund"]]
    counts = Counter(evaluations[key]["bucket"] for key in companies)
    funds = sum(
        1
        for value in evaluations.values()
        if value["fund"] and not value.get("excluded_sector")
    )
    sectors = sum(1 for value in evaluations.values() if value.get("excluded_sector"))
    excellent = sum(
        1
        for key in companies
        if evaluations[key]["quality"] == QUALITY_EXCELLENT
        and evaluations[key]["health"] != QUALITY_WEAK
    )
    traps = sum(
        1 for key in companies if ALERT_VALUE_TRAP in evaluations[key]["alerts"]
    )
    return counts, funds, sectors, excellent, traps, len(companies) + sectors


class GuideWriter:
    """Small row composer that keeps report logic from nesting."""

    NAVY = "0A1B2D"
    PETROL = "003850"
    CARD = "0F2A3F"
    AMBER = "F59C00"
    TEAL = "6DADAA"
    CORAL = "E07068"
    MUTED = "A9B7C2"
    LIGHT = "D8E0E6"

    def __init__(self, sheet):
        self.sheet = sheet
        self.row = 2

    def cell(
        self,
        row,
        column,
        value,
        color=WHITE,
        bold=False,
        size=10,
        fill=None,
        italic=False,
        align="left",
    ):
        cell = self.sheet.cell(row, column, value)
        cell.font = Font(
            name=FONT_NAME, color=color, bold=bold, size=size, italic=italic
        )
        if fill:
            cell.fill = _fill(fill)
        cell.alignment = Alignment(wrap_text=True, vertical=CENTER, horizontal=align)
        return cell

    def _auto_height(self, text, minimum=15):
        self.sheet.row_dimensions[self.row].height = max(
            minimum, 13 * (len(str(text)) // 96 + 1) + 3
        )

    def add_logo(self, assets_dir: Path) -> None:
        try:
            from openpyxl.drawing.image import Image as XLImage

            logo = XLImage(str(assets_dir / "vhc-marca-excel.png"))
            ratio = logo.width / logo.height
            logo.height = 64
            logo.width = int(64 * ratio)
            self.sheet.add_image(logo, "C2")
            self.sheet.row_dimensions[2].height = 50
            self.row = 5
        except (OSError, ValueError):
            self.row = 2

    def title(self, export_name: str, run_date: str, language) -> None:
        self.cell(
            self.row,
            3,
            i18n.term("GUÍA DEL ARCHIVO", language),
            color=self.AMBER,
            bold=True,
        )
        self.row += 1
        self.cell(
            self.row,
            3,
            i18n.term("Screening cuantitativo — VHC Inversiones", language),
            bold=True,
            size=16,
        )
        self.sheet.row_dimensions[self.row].height = 24
        self.row += 1
        self.cell(
            self.row,
            3,
            i18n.text(
                "summary_header",
                language,
                export=export_name,
                run_date=run_date,
                version=SCREENING_VERSION,
            ),
            color=self.MUTED,
            size=9,
            italic=True,
        )
        self.row += 2

    def section(self, number: str, title: str, accent: str) -> None:
        self.cell(
            self.row,
            2,
            number,
            color=self.NAVY,
            bold=True,
            size=13,
            fill=accent,
            align=CENTER,
        )
        self.cell(self.row, 3, f"  {title}", bold=True, size=12, fill=self.PETROL)
        self.sheet.row_dimensions[self.row].height = 22
        self.row += 1

    def text(self, text: str) -> None:
        self.cell(self.row, 3, text, color=self.LIGHT, fill=self.CARD)
        self._auto_height(text)
        self.row += 1

    def labeled(
        self, label: str, text: str, label_fill=None, label_color=WHITE
    ) -> None:
        self.cell(
            self.row,
            2,
            label,
            color=label_color,
            bold=True,
            size=9,
            fill=label_fill or self.CARD,
            align=CENTER,
        )
        self.cell(self.row, 3, text, color=self.LIGHT, fill=self.CARD)
        self._auto_height(text)
        self.row += 1

    def space(self) -> None:
        self.row += 1

    def link(self, label: str, text: str, url: str) -> None:
        self.cell(
            self.row,
            2,
            label,
            color=self.NAVY,
            bold=True,
            size=9,
            fill=self.TEAL,
            align=CENTER,
        )
        link = self.cell(
            self.row, 3, f"  {text} → {url}", color=self.TEAL, bold=True, fill=self.CARD
        )
        link.hyperlink = url
        self.sheet.row_dimensions[self.row].height = 20
        self.row += 1


def _prepare_guide_sheet(workbook, language):
    name = "How to read" if language == i18n.ENGLISH else "Instructivo"
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 2.5
    sheet.column_dimensions["B"].width = 25
    sheet.column_dimensions["C"].width = 104
    sheet.column_dimensions["D"].width = 2.5
    navy_fill = _fill(GuideWriter.NAVY)
    for row in sheet.iter_rows(min_row=1, max_row=189, min_col=1, max_col=4):
        for cell in row:
            cell.fill = navy_fill
    return sheet


def _write_guide(workbook, summary, export_name, run_date, assets_dir, language):
    counts, funds, sectors, excellent, traps, analyzed = summary
    guide = GuideWriter(_prepare_guide_sheet(workbook, language))
    guide.add_logo(assets_dir)
    guide.title(export_name, run_date, language)

    guide.section("01", i18n.term("QUÉ ES ESTA HERRAMIENTA", language), guide.AMBER)
    guide.text(i18n.text("xl_guide_intro", language))
    guide.text(
        i18n.text(
            "xl_guide_run",
            language,
            analyzed=analyzed,
            excellent=excellent,
            deep=counts[BUCKET_DEEP_DIVE],
            watch=counts[BUCKET_WATCHLIST],
            neutral=counts[BUCKET_NEUTRAL],
            discarded=counts[BUCKET_DISCARDED],
            traps=traps,
            funds=funds,
            sectors=sectors,
        )
    )
    guide.space()

    guide.section("02", i18n.term("LAS CATEGORÍAS", language), guide.TEAL)
    categories = (
        (
            "🔬 DEEP DIVE",
            i18n.text("xl_cat_deep", language),
            "A9D08E",
        ),
        (
            "📋 WATCHLIST",
            i18n.text("xl_cat_watch", language),
            "E2EFDA",
        ),
        ("⚪ NEUTRAL", i18n.text("xl_cat_neutral", language), "FFF2CC"),
        (
            i18n.term("❌ DESCARTADA", language),
            i18n.text("xl_cat_discarded", language),
            "FFC7CE",
        ),
        (
            i18n.term("⚙️ OMITIDA", language),
            i18n.text("xl_cat_omitted", language),
            "D6DEE8",
        ),
        (
            i18n.term("FONDO / ETF", language),
            i18n.text("xl_cat_fund", language),
            "E0E0E0",
        ),
    )
    for label, text, color in categories:
        guide.labeled(label, text, color, guide.NAVY)
    guide.space()

    guide.section("03", i18n.term("LOS TRES TRIBUNALES", language), guide.AMBER)
    guide.labeled(
        i18n.term("Calidad", language), i18n.text("xl_gate_quality", language)
    )
    guide.labeled(i18n.term("Salud", language), i18n.text("xl_gate_health", language))
    guide.labeled(i18n.term("Precio", language), i18n.text("xl_gate_price", language))
    guide.space()

    guide.section("04", i18n.term("EDUCATIVO, NO ASESORÍA", language), guide.CORAL)
    guide.text(i18n.text("xl_guide_legal", language))
    guide.space()
    guide.link(
        "YOUTUBE",
        i18n.term("VHC Inversiones en YouTube", language),
        "https://www.youtube.com/@VHCInversiones",
    )
    guide.link(
        "X", i18n.term("VHC Inversiones en X", language), "https://x.com/VHCInversiones"
    )
    guide.link(
        "INSTAGRAM",
        i18n.term("VHC Inversiones en Instagram", language),
        "https://www.instagram.com/vhcinversiones",
    )
    return guide.sheet


def _set_tab_color(sheet: Worksheet, color: str) -> None:
    properties = sheet.sheet_properties
    if properties is None:
        properties = WorksheetProperties()
        sheet.sheet_properties = properties
    properties.tabColor = color


def generate_excel(
    rows: list[Row],
    header_map: HeaderMap,
    evaluations: dict,
    order: list[int],
    output_path: str | Path,
    export_name: str,
    run_date: str,
    assets_dir: str | Path = DEFAULT_ASSETS,
    language: str = i18n.SPANISH,
):
    """Generates the Excel file and returns the executive summary figures."""
    missing_core = [header for header in CORE_COLUMNS if header not in header_map]
    if missing_core:
        joined = ", ".join(missing_core)
        raise ValueError(i18n.text("xl_core_error", language, columns=joined))
    raw, _, _ = make_accessors(header_map)
    missing = [
        header
        for header, _group, calculated in LAYOUT
        if not calculated and header not in header_map
    ]
    offset = int(bool(missing))
    workbook = Workbook()
    screening = workbook.worksheets[0]
    screening.title = "Screening"
    styles = _styles()
    _write_band_headers(screening, offset, language)
    names = _write_column_headers(screening, offset, language)
    _add_header_comments(screening, names, offset, language)
    last_row = _write_data_rows(
        screening, rows, evaluations, order, raw, offset, styles, language
    )
    _mark_missing_columns(screening, missing, names, offset, last_row, language)
    _configure_screening_sheet(screening, offset, names, last_row)
    summary = _summarize(evaluations)
    guide = _write_guide(
        workbook, summary, export_name, run_date, Path(assets_dir), language
    )
    _set_tab_color(screening, "F59C00")
    _set_tab_color(guide, "003850")
    workbook.active = 0
    workbook.save(output_path)
    return summary
