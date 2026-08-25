"""Reading, validation and sector-exclusion tests."""

import builtins
from datetime import date
from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pytest
from openpyxl.worksheet.worksheet import Worksheet
from vhc_screening import i18n
from vhc_screening import io as screening_io
from vhc_screening.io import (
    load_sector_exclusions,
    read_export,
    sector_list_age_warning,
    verify_percentage_scale,
)


def _save_workbook(path, rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.worksheets[0]
    assert isinstance(sheet, Worksheet)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def test_read_export_finds_headers_beyond_the_first_row(tmp_path):
    path = tmp_path / "export.xlsx"
    _save_workbook(
        path,
        [
            ["Export de InvestingPro"],
            ["Name", "Ticker", "Price, Current", "Fair Value"],
            ["Empresa A", "AAA", 10.0, 15.0],
            [None, None, None, None],
        ],
    )

    rows, header_map, missing = read_export(path)

    assert len(rows) == 1
    assert header_map["Price, Current"] == 2
    assert "Fair Value Label (Analyst Targets)" in missing


@pytest.mark.parametrize(
    "rows",
    [
        [["sin", "encabezados"]],
        [["Name", "Ticker"], ["Empresa", "AAA"]],
    ],
)
def test_read_export_rejects_incompatible_files(tmp_path, rows):
    path = tmp_path / "invalido.xlsx"
    _save_workbook(path, rows)

    with pytest.raises(SystemExit, match="ERROR"):
        read_export(path)


@pytest.mark.parametrize(
    ("filename", "content", "message"),
    [
        ("export.xls", b"contenido", "extensión .xlsx"),
        ("vacio.xlsx", b"", "está vacío"),
        ("danado.xlsx", b"no es un zip", "dañado"),
    ],
)
def test_read_export_rejects_empty_or_corrupt_format(
    tmp_path, filename, content, message
):
    path = tmp_path / filename
    path.write_bytes(content)

    with pytest.raises(SystemExit, match=message):
        read_export(path)


def test_read_export_rejects_nonexistent_path(tmp_path):
    with pytest.raises(SystemExit, match="no se encontró"):
        read_export(tmp_path / "ausente.xlsx")


def test_read_export_reports_inspection_error(monkeypatch, tmp_path):
    path = tmp_path / "export.xlsx"
    path.touch()
    real_stat = Path.stat

    def failing_stat(candidate, *args, **kwargs):
        if candidate == path:
            raise OSError("sin permisos")
        return real_stat(candidate, *args, **kwargs)

    monkeypatch.setattr(Path, "stat", failing_stat)
    with pytest.raises(SystemExit, match="no se pudo inspeccionar"):
        read_export(path)


def test_xlsx_limit_is_5_mib_and_accepts_incomplete_real_fixture(
    incomplete_market_view_path,
):
    assert screening_io.MAX_XLSX_FILE_BYTES == 5 * 1024 * 1024
    assert incomplete_market_view_path.stat().st_size < screening_io.MAX_XLSX_FILE_BYTES


def test_read_export_rejects_oversized_external_file(monkeypatch, tmp_path):
    path = tmp_path / "export.xlsx"
    path.write_bytes(b"contenido")
    monkeypatch.setattr(screening_io, "MAX_XLSX_FILE_BYTES", 1)

    with pytest.raises(SystemExit, match="límite de seguridad"):
        read_export(path)


@pytest.mark.parametrize(
    ("constant", "message"),
    [
        ("MAX_XLSX_ARCHIVE_ENTRIES", "demasiados archivos internos"),
        ("MAX_XLSX_MEMBER_BYTES", "archivo interno demasiado grande"),
        ("MAX_XLSX_UNCOMPRESSED_BYTES", "contenido descomprimido"),
    ],
)
def test_read_export_applies_container_limits(monkeypatch, tmp_path, constant, message):
    path = tmp_path / "export.xlsx"
    _save_workbook(path, [["Name", "Ticker"], ["Empresa", "AAA"]])
    monkeypatch.setattr(screening_io, constant, 0)

    with pytest.raises(SystemExit, match=message):
        read_export(path)


def test_read_export_rejects_zip_without_xlsx_structure(tmp_path):
    path = tmp_path / "falso.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr("archivo.txt", "contenido")

    with pytest.raises(SystemExit, match="estructura mínima"):
        read_export(path)


@pytest.mark.parametrize(
    ("constant", "rows", "message"),
    [
        ("MAX_WORKSHEET_ROWS", [["Name", "Ticker"], ["Empresa", "AAA"]], "filas"),
        ("MAX_WORKSHEET_COLUMNS", [["Name", "Ticker"]], "columnas"),
    ],
)
def test_read_export_applies_sheet_limits(
    monkeypatch, tmp_path, constant, rows, message
):
    path = tmp_path / "export.xlsx"
    _save_workbook(path, rows)
    monkeypatch.setattr(screening_io, constant, 1)

    with pytest.raises(SystemExit, match=message):
        read_export(path)


@pytest.mark.parametrize(
    ("constant", "rows", "message"),
    [
        ("MAX_WORKSHEET_ROWS", [["a"], ["b"]], "filas"),
        ("MAX_WORKSHEET_COLUMNS", [["a", "b"]], "columnas"),
    ],
)
def test_limits_apply_to_sheets_without_declared_dimension(
    monkeypatch, constant, rows, message
):
    class UndimensionedSheet:
        max_row = None
        max_column = None

        def iter_rows(self, values_only):
            assert values_only is True
            return iter(rows)

    sheet = UndimensionedSheet()
    monkeypatch.setattr(screening_io, constant, 1)
    with pytest.raises(SystemExit, match=message):
        screening_io._read_sheet_rows(sheet, i18n.SPANISH)


def test_read_export_translates_openpyxl_failures(monkeypatch, tmp_path):
    path = tmp_path / "export.xlsx"
    _save_workbook(path, [["Name", "Ticker"]])

    def fail_load(*_args, **_kwargs):
        raise ValueError("Invalid XML")

    monkeypatch.setattr(screening_io.openpyxl, "load_workbook", fail_load)
    with pytest.raises(SystemExit, match="no se pudo leer el XLSX"):
        read_export(path)


def test_guard_detects_percentages_on_the_wrong_scale():
    header_map = {"Return on Equity": 0}
    decimals = [[0.20] for _ in range(20)]
    percentages = [[20.0] for _ in range(20)]

    assert verify_percentage_scale(decimals, header_map) == []
    assert verify_percentage_scale(percentages, header_map) == ["Return on Equity"]


def test_sector_list_is_mandatory(tmp_path):
    missing_path = tmp_path / "skill" / "references" / "sectores_excluidos.csv"

    with pytest.raises(SystemExit, match="no se encontro"):
        load_sector_exclusions(missing_path)


def test_empty_or_malformed_sector_list_is_rejected(tmp_path):
    csv_path = tmp_path / "skill" / "references" / "sectores_excluidos.csv"
    csv_path.parent.mkdir(parents=True)
    csv_path.write_text(
        "ticker,sector,tipo\n# comentario\nINVALIDO\n", encoding="utf-8"
    )

    with pytest.raises(SystemExit, match="vacia o mal formada"):
        load_sector_exclusions(csv_path)


def test_sector_read_error_aborts(monkeypatch, tmp_path):
    csv_path = tmp_path / "skill" / "references" / "sectores_excluidos.csv"
    csv_path.parent.mkdir(parents=True)
    csv_path.write_text("ticker,sector,tipo\nJPM,banco,Financiero\n", encoding="utf-8")
    real_open = builtins.open

    def failing_open(path, *args, **kwargs):
        if Path(path) == csv_path:
            raise OSError("lectura simulada")
        return real_open(path, *args, **kwargs)

    monkeypatch.setattr(builtins, "open", failing_open)
    with pytest.raises(SystemExit, match="no se pudo leer"):
        load_sector_exclusions(csv_path)


@pytest.mark.parametrize(
    ("cutoff", "expected"),
    [
        ("2026-08-01", None),
        ("2025-01-01", "tiene 588 días de antigüedad"),
        (None, "no declara una fecha de corte"),
        ("2026-99-99", "fecha de corte inválida"),
    ],
)
def test_sector_list_age_warning(tmp_path, cutoff, expected):
    csv_path = tmp_path / "sectores_excluidos.csv"
    heading = f"# corte {cutoff}\n" if cutoff else "# lista sectorial\n"
    csv_path.write_text(f"{heading}ticker,name,sector\n", encoding="utf-8")

    warning = sector_list_age_warning(csv_path, today=date(2026, 8, 12))

    if expected is None:
        assert warning is None
    else:
        assert expected in warning


def test_sector_warning_skips_unreadable_file(tmp_path):
    assert sector_list_age_warning(tmp_path / "ausente.csv") is None
