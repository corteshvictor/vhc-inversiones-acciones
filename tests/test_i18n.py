"""Bilingual output: the language must move text and nothing else."""

import re
from collections import Counter

import openpyxl
import pytest
from vhc_screening import excel_report, i18n
from vhc_screening.constants import (
    ALERT_M_SCORE,
    ALERT_TOO_CHEAP,
    BUCKET_DISCARDED,
    QUALITY_EXCELLENT,
)


def test_terms_translate_only_into_english():
    assert i18n.term(BUCKET_DISCARDED, i18n.SPANISH) == BUCKET_DISCARDED
    assert i18n.term(BUCKET_DISCARDED, i18n.ENGLISH) == "Discarded"
    assert i18n.term(QUALITY_EXCELLENT, i18n.ENGLISH) == "EXCELLENT"
    assert i18n.term("Autodesk Inc", i18n.ENGLISH) == "Autodesk Inc"


def test_alerts_translate_part_by_part():
    joined = f"{ALERT_M_SCORE} · {ALERT_TOO_CHEAP}"
    assert i18n.alert(joined, i18n.SPANISH) == joined
    assert i18n.alert(joined, i18n.ENGLISH) == "M-Score · Too cheap?"


def test_messages_fill_their_placeholders():
    spanish = i18n.text("summary_files", i18n.SPANISH, excel="a", dashboard="b")
    english = i18n.text("summary_files", i18n.ENGLISH, excel="a", dashboard="b")
    assert spanish == "Archivos generados: a · b"
    assert english == "Files generated: a · b"
    assert i18n.text("summary_title", i18n.SPANISH).startswith("══ INFORME")


def test_localized_cells_keep_their_emoji_prefix():
    assert excel_report._localized("🔬 Deep Dive", i18n.SPANISH) == "🔬 Deep Dive"
    assert excel_report._localized("❌ Descartada", i18n.ENGLISH) == "❌ Discarded"
    assert excel_report._localized(BUCKET_DISCARDED, i18n.ENGLISH) == "Discarded"
    assert excel_report._localized(9, i18n.ENGLISH) == 9
    assert excel_report._localized("Autodesk Inc", i18n.ENGLISH) == "Autodesk Inc"


_BACK = {english: spanish for spanish, english in i18n.TERMS.items()}

# Sentences the report builds from a template, so they cannot be reversed by
# a plain lookup: the sector name sits inside them.
_TEMPLATED = ("xl_omitted_cell", "xl_omitted_comment", "xl_fund_cell")


def _from_template(value):
    """Rebuilds the Spanish sentence when an English template produced it."""
    for key in _TEMPLATED:
        spanish, english = i18n.MESSAGES[key]
        pattern = "(.+)".join(re.escape(part) for part in english.split("{sector}"))
        found = re.fullmatch(pattern, value)
        if found is None:
            continue
        if not found.groups():
            return spanish
        sector = found.group(1)
        return spanish.format(sector=_BACK.get(sector, sector))
    return None


def _canonical(value):
    """Maps an English cell back to Spanish so both runs can be compared."""
    if not isinstance(value, str):
        return value
    parts = [part.strip() for part in value.split("·")]
    if len(parts) > 1:
        return " · ".join(_canonical(part) for part in parts)
    if value in _BACK:
        return _BACK[value]
    rebuilt = _from_template(value)
    if rebuilt is not None:
        return rebuilt
    prefix, _, rest = value.partition(" ")
    if rest in _BACK:
        return f"{prefix} {_BACK[rest]}"
    return value


def _run(tmp_path, language, complete_market_view_path, evaluate_row):
    from vhc_screening.cli import run
    from vhc_screening.dashboard import generate_dashboard
    from vhc_screening.io import make_accessors, read_export, verify_percentage_scale

    run(
        read_export=read_export,
        verify_scale=verify_percentage_scale,
        make_accessors=make_accessors,
        evaluate_row=evaluate_row,
        excel_writer=excel_report.generate_excel,
        dashboard_writer=generate_dashboard,
        argv=[
            str(complete_market_view_path),
            "--outdir",
            str(tmp_path),
            "--lang",
            language,
        ],
    )
    workbook = openpyxl.load_workbook(next(tmp_path.glob("Screening_VHC_*.xlsx")))
    sheet = workbook["Screening"]
    verdicts = [
        tuple(
            _canonical(sheet.cell(row, column).value)
            for column in range(1, sheet.max_column + 1)
        )
        for row in range(3, sheet.max_row + 1)
    ]
    names = list(workbook.sheetnames)
    headers = [sheet.cell(2, column).value for column in range(1, 3)]
    workbook.close()
    dashboard = next(tmp_path.glob("Dashboard_VHC_*.html")).read_text(encoding="utf-8")
    return verdicts, names, headers, dashboard


def test_english_run_translates_without_changing_the_classification(
    tmp_path, complete_market_view_path, evaluate_row, capsys
):
    spanish = _run(
        tmp_path / "es", i18n.SPANISH, complete_market_view_path, evaluate_row
    )
    spanish_out = capsys.readouterr().out
    english = _run(
        tmp_path / "en", i18n.ENGLISH, complete_market_view_path, evaluate_row
    )
    english_out = capsys.readouterr().out

    # Every cell of every row, with the English mapped back: bucket, score, the
    # three lenses, each verdict, each signal count and the alerts must match.
    assert spanish[0] == english[0]
    assert len(spanish[0]) == 30
    assert len(spanish[0][0]) == 37

    assert spanish[1] == ["Screening", "Instructivo"]
    assert english[1] == ["Screening", "How to read"]
    assert spanish[2] == ["Balde", "Puntaje (0-10)"]
    assert english[2] == ["Bucket", "Score (0-10)"]

    assert "INFORME EJECUTIVO" in spanish_out
    assert "EXECUTIVE SUMMARY" in english_out
    assert "Descartadas" in spanish_out
    assert "Discarded" in english_out

    assert "setLang('en');" not in spanish[3]
    assert "setLang('en');" in english[3]


def _workbook_shape(path):
    """Everything about the sheet that the language must not touch."""
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Screening"]
    shape = {
        "dimensions": (sheet.max_row, sheet.max_column),
        "merges": sorted(str(area) for area in sheet.merged_cells.ranges),
        "freeze": sheet.freeze_panes,
        "filter": str(sheet.auto_filter.ref),
        "widths": {key: value.width for key, value in sheet.column_dimensions.items()},
        "fills": [
            sheet.cell(row, column).fill.fgColor.rgb
            for row in range(1, sheet.max_row + 1)
            for column in range(1, sheet.max_column + 1)
        ],
        "comments": sum(
            sheet.cell(row, column).comment is not None
            for row in range(1, 4)
            for column in range(1, sheet.max_column + 1)
        ),
    }
    workbook.close()
    return shape


def test_language_moves_text_but_never_formatting(
    tmp_path, complete_market_view_path, evaluate_row, capsys
):
    for language in (i18n.SPANISH, i18n.ENGLISH):
        _run(tmp_path / language, language, complete_market_view_path, evaluate_row)
    capsys.readouterr()

    spanish = _workbook_shape(
        next((tmp_path / i18n.SPANISH).glob("Screening_VHC_*.xlsx"))
    )
    english = _workbook_shape(
        next((tmp_path / i18n.ENGLISH).glob("Screening_VHC_*.xlsx"))
    )

    assert spanish == english
    assert spanish["comments"] == 15


def test_only_the_methods_headers_and_the_guide_change_language(
    tmp_path, complete_market_view_path, evaluate_row, capsys
):
    for language in (i18n.SPANISH, i18n.ENGLISH):
        _run(tmp_path / language, language, complete_market_view_path, evaluate_row)
    capsys.readouterr()

    def sheets(language):
        workbook = openpyxl.load_workbook(
            next((tmp_path / language).glob("Screening_VHC_*.xlsx"))
        )
        screening = workbook["Screening"]
        headers = [
            screening.cell(2, column).value
            for column in range(1, screening.max_column + 1)
        ]
        guide = workbook.worksheets[1]
        prose = " ".join(
            str(cell.value)
            for row in guide.iter_rows()
            for cell in row
            if isinstance(cell.value, str)
        )
        workbook.close()
        return headers, prose

    spanish_headers, spanish_guide = sheets(i18n.SPANISH)
    english_headers, english_guide = sheets(i18n.ENGLISH)

    # Every header the catalog knows must have moved; the InvestingPro column
    # names are already English and must stay exactly as the export wrote them.
    translated = [
        (spanish, english)
        for spanish, english in zip(spanish_headers, english_headers, strict=True)
        if spanish != english
    ]
    assert len(translated) == 9
    assert all(
        i18n.TERMS.get(str(spanish)) == english for spanish, english in translated
    )

    for phrase in ("QUÉ ES ESTA HERRAMIENTA", "LAS CATEGORÍAS", "EDUCATIVO"):
        assert phrase in spanish_guide
        assert phrase not in english_guide
    for phrase in ("WHAT THIS TOOL IS", "THE CATEGORIES", "EDUCATIONAL"):
        assert phrase in english_guide
        assert phrase not in spanish_guide


def test_reader_errors_follow_the_requested_language(tmp_path):
    from vhc_screening.io import read_export

    missing = tmp_path / "nope.xlsx"

    with pytest.raises(SystemExit) as spanish:
        read_export(missing, i18n.SPANISH)
    with pytest.raises(SystemExit) as english:
        read_export(missing, i18n.ENGLISH)

    assert "no se encontró el archivo XLSX" in str(spanish.value)
    assert "the XLSX file was not found" in str(english.value)


def test_dashboard_core_error_follows_the_requested_language(tmp_path):
    from vhc_screening.dashboard import generate_dashboard

    def build(language):
        generate_dashboard(
            {},
            [],
            tmp_path / f"{language}.html",
            "fixture.xlsx",
            "2026-08-22",
            Counter(),
            0,
            0,
            0,
            0,
            missing=["Beneish M-Score"],
            language=language,
        )

    with pytest.raises(ValueError, match="faltan columnas núcleo"):
        build(i18n.SPANISH)
    with pytest.raises(ValueError, match="core columns are missing"):
        build(i18n.ENGLISH)


def test_dashboard_is_born_in_the_requested_language(
    tmp_path, complete_market_view_path, evaluate_row, capsys
):
    for language in (i18n.SPANISH, i18n.ENGLISH):
        _run(tmp_path / language, language, complete_market_view_path, evaluate_row)
    capsys.readouterr()

    def markup(language):
        return next((tmp_path / language).glob("Dashboard_VHC_*.html")).read_text(
            encoding="utf-8"
        )

    spanish, english = markup(i18n.SPANISH), markup(i18n.ENGLISH)

    assert '<html lang="es"' in spanish
    assert '<html lang="en"' in english
    assert "<title>Screening Cuantitativo" in spanish
    assert "<title>Quantitative Screening" in english
    assert 'data-lang="es" aria-selected="true"' in spanish
    assert 'data-lang="en" aria-selected="true"' in english
    assert "setLang('en');" not in spanish
    assert "setLang('en');" in english


def _placeholders(text):
    return set(re.findall(r"\{(\w+)\}", text))


def test_the_catalog_is_internally_consistent():
    """A future translation cannot drop a placeholder or collide with another."""
    for key, (spanish, english) in i18n.MESSAGES.items():
        assert spanish, f"{key} has no Spanish text"
        assert english, f"{key} has no English text"
        assert spanish != english, f"{key} was never translated"
        assert _placeholders(spanish) == _placeholders(english), (
            f"{key} has mismatched placeholders"
        )

    # No two Spanish values may share an English rendering: the reverse mapping
    # used to compare both runs would become ambiguous.
    assert len(set(i18n.TERMS.values())) == len(i18n.TERMS)

    spanish_words = re.compile(
        r"\b(el|la|los|las|del|que|para|con|una|señal|señales|salud|calidad|"
        r"precio|empresa|empresas|umbral|umbrales)\b",
        re.IGNORECASE,
    )
    for key, (_spanish, english) in i18n.MESSAGES.items():
        assert not spanish_words.search(english), f"{key} leaks Spanish"
    for value, english in i18n.TERMS.items():
        assert not spanish_words.search(english), f"{value} leaks Spanish"


def test_reader_errors_reach_english_through_run(tmp_path, capsys):
    """The injectable contract must carry the language, not just main()."""
    from vhc_screening.cli import run
    from vhc_screening.dashboard import generate_dashboard
    from vhc_screening.io import make_accessors, read_export, verify_percentage_scale

    def attempt(language):
        return run(
            read_export=read_export,
            verify_scale=verify_percentage_scale,
            make_accessors=make_accessors,
            evaluate_row=lambda *_a: {},
            excel_writer=excel_report.generate_excel,
            dashboard_writer=generate_dashboard,
            argv=[
                str(tmp_path / "absent.xlsx"),
                "--outdir",
                str(tmp_path),
                "--lang",
                language,
            ],
        )

    with pytest.raises(SystemExit) as spanish:
        attempt(i18n.SPANISH)
    with pytest.raises(SystemExit) as english:
        attempt(i18n.ENGLISH)

    assert "no se encontró el archivo XLSX" in str(spanish.value)
    assert "the XLSX file was not found" in str(english.value)
    capsys.readouterr()


def test_raw_export_values_are_never_translated(make_case, evaluate_row, tmp_path):
    """A source cell that happens to read like a method term stays untouched."""
    assert "Calidad" in i18n.TERMS  # would be turned into "Quality" if localized

    row, header_map, raw, number, percentage = make_case(
        {"Overall Health Label": "Calidad"}
    )
    evaluation = evaluate_row(row, raw, number, percentage)
    output = tmp_path / "raw.xlsx"
    excel_report.generate_excel(
        [row],
        header_map,
        {0: evaluation},
        [0],
        output,
        "fixture.xlsx",
        "2026-08-23",
        language=i18n.ENGLISH,
    )

    workbook = openpyxl.load_workbook(output)
    sheet = workbook["Screening"]
    headers = {sheet.cell(2, column).value: column for column in range(1, 38)}
    assert sheet.cell(3, headers["Overall Health Label"]).value == "Calidad"
    assert str(sheet.cell(3, headers["Bucket"]).value).endswith("Deep Dive")
    workbook.close()
