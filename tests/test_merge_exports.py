"""Tests for the tool that joins several exports into one universe.

The behaviour worth guarding is not the merging — any script concatenates rows —
but the continuity check. An export series assembled by lowering the market
capitalization filter loses companies silently when the filter is rounded down,
and no later reader of the file can tell. These tests pin the rule that catches
it: every boundary must repeat a company, and the resume value always rounds up.
"""

import os
import runpy
import sys
from unittest import mock

import openpyxl
import pytest
from merge_exports import (
    CLASSIFIER_ROW_LIMIT,
    FULL_TICKER,
    MARKET_CAP,
    NAME,
    TICKER,
    Batch,
    Gap,
    _number,
    _text,
    find_gaps,
    main,
    merge,
    read_batch,
    require_matching_columns,
    resolve_paths,
    write_sheet,
)

HEADER = [NAME, TICKER, MARKET_CAP]


def _export(path, rows, *, header=None, preamble=0):
    """Writes a sheet shaped like an InvestingPro export."""
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    for _ in range(preamble):
        sheet.append(["Untitled Screener"])
    sheet.append(header if header is not None else HEADER)
    for row in rows:
        sheet.append(list(row))
    book.save(path)
    return path


def _slice(tmp_path, name, tickers, caps, **kwargs):
    rows = [(f"{t} Inc", t, c) for t, c in zip(tickers, caps, strict=True)]
    return _export(tmp_path / name, rows, **kwargs)


def test_text_collapses_whitespace_and_survives_empty_cells():
    assert _text("  Deckers   Outdoor \n Corp ") == "Deckers Outdoor Corp"
    assert _text(None) == ""


def test_number_accepts_only_real_numbers():
    assert _number(292_394_410) == 292_394_410.0
    assert _number(1.5) == 1.5
    assert _number("292394410") is None
    assert _number(None) is None
    # A bool is an int in Python, and a market cap of True would sort as 1.0.
    assert _number(True) is None


def test_read_batch_finds_the_header_below_a_preamble(tmp_path):
    path = _slice(tmp_path, "a.xlsx", ["NVDA", "AAPL"], [5e12, 4e12], preamble=7)
    batch = read_batch(str(path))

    assert batch.tickers == ["NVDA", "AAPL"]
    assert batch.top == 5e12
    assert batch.bottom == 4e12
    assert batch.name == "a.xlsx"


def test_read_batch_skips_rows_with_no_company_name(tmp_path):
    path = _export(tmp_path / "a.xlsx", [("NVDA Inc", "NVDA", 5e12), (None, "", None)])
    assert len(read_batch(str(path)).rows) == 1


def test_read_batch_refuses_an_export_with_no_capitalization(tmp_path):
    """Without that column every batch sorts as zero, so the ordering and the
    boundary check both become theatre: they would report continuity over a
    sequence nobody established."""
    path = _export(tmp_path / "a.xlsx", [("NVDA Inc", "NVDA")], header=[NAME, TICKER])

    with pytest.raises(SystemExit, match="no trae la columna"):
        read_batch(str(path))


def test_read_batch_refuses_a_file_with_no_header_row(tmp_path):
    path = _export(tmp_path / "a.xlsx", [], header=["Something", "Else"])
    with pytest.raises(SystemExit, match="fila de encabezados"):
        read_batch(str(path))


def test_read_batch_refuses_a_workbook_with_no_sheet(tmp_path, monkeypatch):
    path = _slice(tmp_path, "a.xlsx", ["NVDA"], [5e12])

    class Sheetless:
        active = None

        def close(self):
            return None

    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: Sheetless())
    with pytest.raises(SystemExit, match="ninguna hoja"):
        read_batch(str(path))


def test_a_boundary_that_repeats_a_company_reports_no_gap(tmp_path):
    first = read_batch(str(_slice(tmp_path, "a.xlsx", ["NVDA", "PTCT"], [5e12, 5.9e9])))
    second = read_batch(
        str(_slice(tmp_path, "b.xlsx", ["PTCT", "UVE"], [5.9e9, 1.2e9]))
    )

    assert find_gaps([first, second]) == []


def test_a_boundary_with_no_overlap_names_the_range_that_was_lost(tmp_path):
    first = read_batch(
        str(_slice(tmp_path, "a.xlsx", ["NVDA", "UVE"], [5e12, 1_213_403_478]))
    )
    second = read_batch(
        str(_slice(tmp_path, "b.xlsx", ["ODTX", "RCMT"], [1_205_149_243, 2.9e8]))
    )

    gaps = find_gaps([first, second])

    assert len(gaps) == 1
    assert gaps[0].after == "a.xlsx"
    assert gaps[0].before == "b.xlsx"
    assert gaps[0].lost_from == 1_205_149_243
    assert gaps[0].lost_to == 1_213_403_478


def test_the_resume_value_always_rounds_up():
    """Rounding down is what opens the gap: the boundary company then falls
    outside both exports, and nothing in the file records its absence."""
    gap = Gap(after="a", before="b", lost_from=1_205_149_243, lost_to=1_213_403_478)

    assert gap.resume_at == 1213.41  # 1,213.403478 M, never 1213.40
    assert gap.resume_at * 1e6 > gap.lost_to

    exact = Gap(after="a", before="b", lost_from=0.0, lost_to=100_000_000)
    assert exact.resume_at == 100.0


def test_merge_keeps_the_first_appearance_of_every_ticker(tmp_path):
    first = read_batch(str(_slice(tmp_path, "a.xlsx", ["NVDA", "PTCT"], [5e12, 5.9e9])))
    second = read_batch(
        str(_slice(tmp_path, "b.xlsx", ["PTCT", "UVE"], [5.9e9, 1.2e9]))
    )

    header, rows, repeated = merge([first, second])

    assert header == HEADER
    assert [row[1] for row in rows] == ["NVDA", "PTCT", "UVE"]
    assert repeated == ["PTCT"]


def test_merge_sorts_the_final_rows_by_capitalization(tmp_path):
    caps = [float(value) * 1e6 for value in range(102, 0, -1)]
    caps[50], caps[51] = caps[51], caps[50]
    tickers = [f"T{value}" for value in range(102, 0, -1)]
    path = _slice(
        tmp_path,
        "a.xlsx",
        [*tickers, "NONE"],
        [*caps, None],
    )

    _, rows, _ = merge([read_batch(str(path))])

    assert [row[2] for row in rows[:-1]] == sorted(caps, reverse=True)
    assert rows[-1][1:] == ["NONE", None]


def test_merge_keeps_rows_that_carry_no_ticker(tmp_path):
    path = _export(tmp_path / "a.xlsx", [("Sin ticker", None, 1e9)])
    _, rows, repeated = merge([read_batch(str(path))])

    assert len(rows) == 1
    assert repeated == []


def test_write_sheet_puts_one_header_above_the_rows(tmp_path):
    output = tmp_path / "universe.xlsx"
    write_sheet(HEADER, [["NVDA Inc", "NVDA", 5e12]], str(output))

    book = openpyxl.load_workbook(output)
    sheet = book.active
    assert sheet is not None
    assert [cell.value for cell in sheet[1]] == HEADER
    assert sheet.max_row == 2
    book.close()


def test_resolve_paths_expands_a_pattern_and_drops_excel_lock_files(tmp_path):
    _slice(tmp_path, "one.xlsx", ["A"], [1e9])
    _slice(tmp_path, "two.xlsx", ["B"], [2e9])
    (tmp_path / "~$one.xlsx").write_bytes(b"lock")

    found = resolve_paths([str(tmp_path / "*.xlsx")])

    # os.path.basename, not a split on "/": Windows separates with a backslash.
    assert [os.path.basename(p) for p in found] == ["one.xlsx", "two.xlsx"]


def test_resolve_paths_takes_the_list_a_shell_already_expanded(tmp_path):
    """Quoting the pattern hands one string; leaving it unquoted hands one
    argument per file. Both forms have to work, and overlap must not duplicate."""
    first = str(_slice(tmp_path, "one.xlsx", ["A"], [1e9]))
    second = str(_slice(tmp_path, "two.xlsx", ["B"], [2e9]))

    assert resolve_paths([first, second]) == sorted([first, second])
    assert resolve_paths([first, str(tmp_path / "*.xlsx")]) == sorted([first, second])


def test_main_merges_orders_by_capitalization_and_reports_no_gaps(tmp_path, capsys):
    # Named so the alphabetical order contradicts the real one: only the data
    # settles the sequence.
    _slice(tmp_path, "z-first.xlsx", ["NVDA", "PTCT"], [5e12, 5.9e9])
    _slice(tmp_path, "a-second.xlsx", ["PTCT", "UVE"], [5.9e9, 1.2e9])
    output = tmp_path / "universe.xlsx"

    code = main([str(tmp_path / "*.xlsx"), "-o", str(output)])
    printed = capsys.readouterr().out

    assert code == 0
    assert "Sin huecos" in printed
    assert printed.index("z-first.xlsx") < printed.index("a-second.xlsx")
    assert "empresas únicas : 3" in printed
    assert output.exists()


def test_main_reports_a_gap_and_exits_non_zero(tmp_path, capsys):
    _slice(tmp_path, "a.xlsx", ["NVDA", "UVE"], [5e12, 1_213_403_478])
    _slice(tmp_path, "b.xlsx", ["ODTX", "RCMT"], [1_205_149_243, 2.9e8])

    code = main([str(tmp_path / "*.xlsx"), "-o", str(tmp_path / "u.xlsx")])
    printed = capsys.readouterr().out

    assert code == 1
    assert "HUECO" in printed
    assert "1213.41" in printed
    assert "million" in printed  # the four boxes of the web filter
    assert '"$lte": 1213.41' in printed  # the API form


def test_main_writes_beside_the_exports_when_no_output_is_given(tmp_path, capsys):
    """The result belongs with its sources, not in whatever directory the
    command happened to run from."""
    _slice(tmp_path, "a.xlsx", ["NVDA"], [5e12])

    main([str(tmp_path / "*.xlsx")])
    capsys.readouterr()

    assert (tmp_path / "universe.xlsx").exists()


def test_main_says_nothing_about_continuity_for_a_single_export(tmp_path, capsys):
    """One export has no boundary, so there is nothing to claim either way."""
    _slice(tmp_path, "a.xlsx", ["NVDA", "PTCT"], [5e12, 5.9e9])

    code = main([str(tmp_path / "*.xlsx"), "-o", str(tmp_path / "u.xlsx")])
    printed = capsys.readouterr().out

    assert code == 0
    assert "CONTINUIDAD" not in printed
    assert "Sin huecos" not in printed


def test_main_tells_the_user_what_to_type_for_the_next_export(tmp_path, capsys):
    caps = [5e12 - index for index in range(1000)]
    _slice(tmp_path, "a.xlsx", [f"T{i}" for i in range(1000)], caps)

    main([str(tmp_path / "*.xlsx"), "-o", str(tmp_path / "u.xlsx")])
    printed = capsys.readouterr().out

    assert "SIGUIENTE TANDA" in printed
    assert "less than or equal to" in printed


def test_main_stops_asking_once_an_export_comes_back_short(tmp_path, capsys):
    _slice(tmp_path, "a.xlsx", ["NVDA", "PTCT"], [5e12, 5.9e9])

    main([str(tmp_path / "*.xlsx"), "-o", str(tmp_path / "u.xlsx")])
    printed = capsys.readouterr().out

    assert "universo está completo" in printed
    assert "SIGUIENTE TANDA" not in printed


def test_main_refuses_a_pattern_that_matches_nothing(tmp_path):
    with pytest.raises(SystemExit, match="ningún archivo coincide"):
        main([str(tmp_path / "nada-*.xlsx")])


def test_an_empty_batch_reports_zero_rather_than_raising():
    batch = Batch(
        path="a.xlsx",
        header=HEADER,
        columns={},
        rows=[],
        tickers=[],
        identities=[],
        caps=[],
    )
    assert batch.top == 0.0
    assert batch.bottom == 0.0
    assert batch.top_ticker == ""
    assert batch.bottom_ticker == ""


def test_the_tool_runs_as_a_command(project_root, tmp_path, capsys):
    """The `if __name__` guard: executed directly the tool must exit with the
    code it returns, so a gap makes a shell script or a CI step fail."""
    _slice(tmp_path, "a.xlsx", ["NVDA", "PTCT"], [5e12, 5.9e9])
    argv = [
        "merge_exports.py",
        str(tmp_path / "*.xlsx"),
        "-o",
        str(tmp_path / "u.xlsx"),
    ]

    with (
        mock.patch.object(sys, "argv", argv),
        pytest.raises(SystemExit) as exit_info,
    ):
        runpy.run_path(
            str(project_root / "tools" / "merge_exports.py"), run_name="__main__"
        )

    capsys.readouterr()
    assert exit_info.value.code == 0


def test_columns_are_read_by_name_not_by_position(tmp_path):
    """Two exports may list the same columns in a different order. Lining them
    up by position slides every value one cell sideways and says nothing: the
    ticker lands under the company name and the sheet still looks plausible."""
    first = _export(tmp_path / "a.xlsx", [("Nvidia", "NVDA", 5e12)])
    second = _export(
        tmp_path / "b.xlsx",
        [("AAPL", "Apple", 4e12)],
        header=[TICKER, NAME, MARKET_CAP],
    )

    header, rows, _ = merge([read_batch(str(first)), read_batch(str(second))])

    assert header == HEADER
    assert rows[1][0] == "Apple"  # under Name, not the ticker
    assert rows[1][1] == "AAPL"


def test_a_batch_with_different_columns_stops_the_merge(tmp_path):
    """A screener reconfigured between downloads yields a different column set,
    and a sheet built from both holds blanks where the method expects figures."""
    first = _export(tmp_path / "a.xlsx", [("Nvidia", "NVDA", 5e12)])
    wider = _export(
        tmp_path / "b.xlsx",
        [("Meta", "META", 1e12, 20)],
        header=[NAME, TICKER, MARKET_CAP, "P/E Ratio"],
    )

    # Read both outside the block: with the reads inside it, a failure there
    # would satisfy pytest.raises and the test would pass without merge() ever
    # having refused anything.
    batches = [read_batch(str(first)), read_batch(str(wider))]

    with pytest.raises(SystemExit) as failure:
        merge(batches)

    assert "trae de más: P/E Ratio" in str(failure.value)
    assert "b.xlsx" in str(failure.value)


def test_a_batch_missing_a_column_stops_the_merge(tmp_path):
    wide = Batch(
        path="a.xlsx",
        header=[NAME, TICKER, MARKET_CAP],
        columns={NAME: 0, TICKER: 1, MARKET_CAP: 2},
        rows=[],
        tickers=[],
        identities=[],
        caps=[1.0],
    )
    narrow = Batch(
        path="b.xlsx",
        header=[NAME, TICKER],
        columns={NAME: 0, TICKER: 1},
        rows=[],
        tickers=[],
        identities=[],
        caps=[1.0],
    )

    with pytest.raises(SystemExit) as failure:
        require_matching_columns([wide, narrow])

    assert f"le faltan: {MARKET_CAP}" in str(failure.value)


def test_the_mirrored_row_limit_matches_the_classifier():
    """merge_exports keeps its own copy of the ceiling so it runs standalone,
    without importing the package. This is what keeps the copy honest."""
    from vhc_screening import io as screening_io

    assert CLASSIFIER_ROW_LIMIT == screening_io.MAX_WORKSHEET_ROWS


def test_the_ceiling_counts_the_header_row(tmp_path, capsys, monkeypatch):
    """The classifier limits the whole sheet, header included, so a limit of 3
    fits two companies and not three. Off by one here writes a file that the
    screening then refuses, which is the worst of both."""
    monkeypatch.setattr("merge_exports.CLASSIFIER_ROW_LIMIT", 3)
    _slice(tmp_path, "a.xlsx", ["AAA", "BBB"], [5e12, 4e12])
    output = tmp_path / "u.xlsx"

    assert main([str(tmp_path / "a.xlsx"), "-o", str(output)]) == 0
    capsys.readouterr()
    assert output.exists()


def test_a_sheet_past_the_ceiling_is_refused_without_writing(
    tmp_path, capsys, monkeypatch
):
    monkeypatch.setattr("merge_exports.CLASSIFIER_ROW_LIMIT", 3)
    _slice(tmp_path, "a.xlsx", ["AAA", "BBB", "CCC"], [5e12, 4e12, 3e12])
    output = tmp_path / "u.xlsx"

    with pytest.raises(SystemExit) as failure:
        main([str(tmp_path / "a.xlsx"), "-o", str(output)])

    capsys.readouterr()
    assert "4 filas" in str(failure.value)
    assert not output.exists(), "no debe escribir una hoja que el screening rechaza"


def test_a_repeated_header_is_reported_rather_than_refused(tmp_path, capsys):
    """A real InvestingPro export repeats 'Full Ticker', so refusing duplicates
    would refuse every genuine file. The last one wins, which is what the
    classifier reads as well."""
    path = _export(
        tmp_path / "a.xlsx",
        [("NVDA Inc", "NVDA", 5e12, "NASDAQGS:NVDA", "NasdaqGS:NVDA")],
        header=[NAME, TICKER, MARKET_CAP, "Full Ticker", "Full Ticker"],
    )

    batch = read_batch(str(path))
    printed = capsys.readouterr().out

    assert "repite 'Full Ticker'" in printed
    assert batch.columns["Full Ticker"] == 4
    assert len(batch.rows) == 1


@pytest.mark.parametrize("repeated", [TICKER, NAME, MARKET_CAP])
def test_any_header_other_than_full_ticker_stops_the_read(tmp_path, repeated):
    """Full Ticker is tolerated because InvestingPro really does repeat it. The
    exception ends there: a second Ticker or Market Cap column takes over the
    first, and nothing downstream can tell which one it read."""
    header = [NAME, TICKER, MARKET_CAP, repeated]
    path = _export(
        tmp_path / "a.xlsx", [("Nvidia", "NVDA", 5e12, "MAL")], header=header
    )

    with pytest.raises(SystemExit) as failure:
        read_batch(str(path))

    assert f"repite la columna '{repeated}'" in str(failure.value)


def test_blank_capitalizations_at_the_edges_do_not_define_the_range(tmp_path, capsys):
    rows = [
        ("Sin capitalización", "NONE1", None),
        ("Mayor", "TOP", 5e12),
        ("Menor", "LOW", 1e9),
        ("Otra sin capitalización", "NONE2", None),
    ]
    path = _export(tmp_path / "a.xlsx", rows)

    batch = read_batch(str(path))

    assert batch.top_ticker == "TOP"
    assert batch.bottom_ticker == "LOW"
    assert "2 filas sin" in capsys.readouterr().out


def test_a_gap_in_the_middle_is_kept_and_only_reported(tmp_path, capsys):
    """Capitalization is optional for the method, so an interior blank merges
    like any other row. Only the ordering ignores it."""
    path = _export(
        tmp_path / "a.xlsx",
        [("Uno", "AAA", 5e12), ("Medio", "BBB", None), ("Fin", "CCC", 1e9)],
    )

    batch = read_batch(str(path))

    assert len(batch.rows) == 3
    assert "1 filas sin" in capsys.readouterr().out


def test_an_export_with_no_companies_says_so(tmp_path):
    path = _export(tmp_path / "a.xlsx", [])

    with pytest.raises(SystemExit, match="ninguna empresa"):
        read_batch(str(path))


def test_a_blank_cell_in_the_header_row_is_ignored(tmp_path):
    """Exports carry empty cells between column groups; they name no column."""
    path = _export(
        tmp_path / "a.xlsx",
        [("Nvidia", "NVDA", 5e12, None)],
        header=[NAME, TICKER, MARKET_CAP, None],
    )

    batch = read_batch(str(path))

    assert MARKET_CAP in batch.columns
    assert "" not in batch.columns


def test_a_row_shorter_than_the_header_yields_blanks_not_an_error():
    """A sheet may end its rows early when the trailing cells are empty. Those
    columns come back blank rather than raising, which is what the classifier
    does with them too."""
    batch = Batch(
        path="a.xlsx",
        header=HEADER,
        columns={NAME: 0, TICKER: 1, MARKET_CAP: 2},
        rows=[["Nvidia", "NVDA"]],  # the capitalization cell is simply absent
        tickers=["NVDA"],
        identities=["nvda\x00nvidia"],
        caps=[5e12],
    )

    _, rows, _ = merge([batch])

    assert rows == [["Nvidia", "NVDA", None]]


def test_an_isolated_row_out_of_order_is_tolerated_and_sorted(tmp_path, capsys):
    """A real export arrives with rows out of order. In one downloaded on 18
    September 2026, the slice from 289.84 M down, row 117 holds CAVG
    (254.46 M) right after AMDO.F (252.60 M). An earlier rule refused any batch
    that rose between neighbors, so that single row threw the whole file away.
    The tool tolerates isolated export noise and sorts it instead; the edges
    come from the values, not the positions."""
    caps = [float(value) * 1e6 for value in range(1_000, 0, -1)]
    caps[115], caps[116] = caps[116], caps[115]
    rows = [(f"Empresa {i}", f"T{i}", cap) for i, cap in enumerate(caps)]
    path = _export(tmp_path / "a.xlsx", rows)

    batch = read_batch(str(path))

    assert len(batch.rows) == len(caps)
    assert batch.top == 1_000e6
    assert batch.bottom == 1e6
    assert batch.top_ticker == "T0"
    assert batch.bottom_ticker == "T999"
    assert "1 de 999 pares consecutivos sube" in capsys.readouterr().out


def _descending_with_rises(size, rises):
    """Capitalizations from `size` million down, with `rises` adjacent swaps.

    The swaps sit ten rows apart, so each one adds exactly one rise and no two
    of them touch.
    """
    caps = [float(value) * 1e6 for value in range(size, 0, -1)]
    for at in range(0, 10 * rises, 10):
        caps[at], caps[at + 1] = caps[at + 1], caps[at]
    return caps


@pytest.mark.parametrize(
    ("size", "rises"),
    [
        (1_000, 9),  # 9 of 999 comparisons, 0.90%: inside the 1%
        (3, 1),  # 1 of 2, 50%: a single rise passes whatever the ratio
    ],
    ids=["nine-rises-in-a-full-slice", "one-rise-in-a-small-file"],
)
def test_the_tolerance_accepts_up_to_its_limit(tmp_path, capsys, size, rises):
    caps = _descending_with_rises(size, rises)
    path = _slice(tmp_path, "a.xlsx", [f"T{i}" for i in range(size)], caps)

    batch = read_batch(str(path))

    assert len(batch.rows) == size
    assert f"{rises} de {size - 1} pares consecutivos" in capsys.readouterr().out


def test_the_tolerance_refuses_one_rise_beyond_its_limit(tmp_path):
    """Ten rises in a full slice are 10 of 999, just over 1%. Together with the
    nine accepted above, this pins the threshold itself: raising it, or
    dropping the single-rise allowance, makes one of these tests fail."""
    caps = _descending_with_rises(1_000, 10)
    path = _slice(tmp_path, "a.xlsx", [f"T{i}" for i in range(1_000)], caps)

    with pytest.raises(SystemExit) as stopped:
        read_batch(str(path))

    message = str(stopped.value)
    assert "10 de 999 pares consecutivos suben" in message
    assert "hasta el 1%" in message


def test_equal_capitalizations_side_by_side_are_not_a_rise(tmp_path, capsys):
    """Ties are an order, not a break. The exports of 18 September 2026 carry
    them: PTOR and CAII both stand at exactly 338,552,368. Counted as rises,
    every tie would eat into the tolerance meant for rows out of place."""
    path = _slice(
        tmp_path,
        "a.xlsx",
        ["TOP", "PTOR", "CAII", "LOW"],
        [500e6, 338_552_368, 338_552_368, 100e6],
    )

    read_batch(str(path))

    assert "pares consecutivos" not in capsys.readouterr().out


def test_an_export_clearly_not_sorted_by_market_cap_stops_the_read(tmp_path):
    caps = [float(value) * 1e6 for value in range(1, 101)]
    path = _slice(
        tmp_path,
        "sorted-by-another-column.xlsx",
        [f"T{value}" for value in range(1, 101)],
        caps,
    )

    with pytest.raises(SystemExit) as stopped:
        read_batch(str(path))

    message = str(stopped.value)
    assert "no conserva el orden descendente" in message
    assert "99 de 99 pares consecutivos suben" in message
    assert "InvestingPro" in message
    assert "de mayor a menor" in message


def test_an_export_with_no_numeric_capitalizations_stops_the_read(tmp_path):
    path = _export(tmp_path / "a.xlsx", [("Uno", "AAA", None)])

    with pytest.raises(SystemExit, match="ningún valor numérico"):
        read_batch(str(path))


def test_two_companies_sharing_a_ticker_both_survive(tmp_path):
    """A ticker does not identify a company. In an August 2026 export, FLYX is
    both Flyexclusive Inc on NYSE and FlyE Group Inc on NYSEAM. Keyed on the
    ticker, the second is dropped and the merged sheet gives no sign of it."""
    path = _export(
        tmp_path / "a.xlsx",
        [
            ("Flyexclusive Inc", "FLYX", 5e8, "NYSE:FLYX"),
            ("FlyE Group Inc", "FLYX", 4e8, "NYSEAM:FLYX"),
        ],
        header=[NAME, TICKER, MARKET_CAP, FULL_TICKER],
    )

    _, rows, repeated = merge([read_batch(str(path))])

    assert len(rows) == 2
    assert repeated == []


def test_the_same_company_across_two_batches_is_still_one(tmp_path):
    """The boundary duplicate must keep collapsing: it is the same listing."""
    first = _export(
        tmp_path / "a.xlsx",
        [
            ("Nvidia", "NVDA", 5e12, "NASDAQGS:NVDA"),
            ("PTC", "PTCT", 5.9e9, "NASDAQGS:PTCT"),
        ],
        header=[NAME, TICKER, MARKET_CAP, FULL_TICKER],
    )
    second = _export(
        tmp_path / "b.xlsx",
        # The exchange spelled differently, as the two Full Ticker columns do.
        [
            ("PTC", "PTCT", 5.9e9, "NasdaqGS:PTCT"),
            ("Universal", "UVE", 1.2e9, "NYSE:UVE"),
        ],
        header=[NAME, TICKER, MARKET_CAP, FULL_TICKER],
    )

    batches = [read_batch(str(first)), read_batch(str(second))]

    assert find_gaps(batches) == []
    _, rows, repeated = merge(batches)
    assert len(rows) == 3
    assert repeated == ["PTCT"]


def test_without_full_ticker_identity_falls_back_to_ticker_and_name(tmp_path):
    """`Full Ticker` is optional for the method, so its absence cannot stop the
    merge. Ticker plus name still tells two companies apart."""
    path = _export(
        tmp_path / "a.xlsx",
        [("Flyexclusive Inc", "FLYX", 5e8), ("FlyE Group Inc", "FLYX", 4e8)],
    )

    _, rows, repeated = merge([read_batch(str(path))])

    assert len(rows) == 2
    assert repeated == []


def test_a_boundary_of_two_different_companies_is_not_continuity(tmp_path):
    """Matching tickers across a boundary prove nothing if the companies differ:
    the range between them was never covered."""
    first = _export(
        tmp_path / "a.xlsx",
        [
            ("Nvidia", "NVDA", 5e12, "NASDAQGS:NVDA"),
            ("Flyexclusive", "FLYX", 5e8, "NYSE:FLYX"),
        ],
        header=[NAME, TICKER, MARKET_CAP, FULL_TICKER],
    )
    second = _export(
        tmp_path / "b.xlsx",
        [("FlyE Group", "FLYX", 4e8, "NYSEAM:FLYX"), ("Otra", "ZZZ", 1e8, "NYSE:ZZZ")],
        header=[NAME, TICKER, MARKET_CAP, FULL_TICKER],
    )

    gaps = find_gaps([read_batch(str(first)), read_batch(str(second))])

    assert len(gaps) == 1, "same ticker, different companies: no continuity proved"


def test_a_third_full_ticker_column_stops_the_read(tmp_path):
    """Two is what InvestingPro writes, and the last one decides identity. A
    third makes that choice unpredictable, and identity is what keeps two
    companies sharing a ticker apart."""
    path = _export(
        tmp_path / "a.xlsx",
        [("Nvidia", "NVDA", 5e12, "A", "B", "C")],
        header=[NAME, TICKER, MARKET_CAP, FULL_TICKER, FULL_TICKER, FULL_TICKER],
    )

    with pytest.raises(SystemExit) as failure:
        read_batch(str(path))

    assert "3 veces" in str(failure.value)
    assert "4, 5, 6" in str(failure.value)
