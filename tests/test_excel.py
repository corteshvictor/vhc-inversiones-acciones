"""Excel output tests."""

from collections import Counter

import openpyxl
import pytest
from openpyxl.worksheet.worksheet import Worksheet
from vhc_screening.excel_report import generate_excel


def _column_for_header(sheet: Worksheet, header: str) -> int:
    for cell in sheet[2]:
        if cell.value == header:
            column = cell.column
            assert isinstance(column, int)
            return column
    raise AssertionError(f"Column not found: {header!r}")


def test_excel_covers_deep_dive_watchlist_fund_and_sector(report_cases, tmp_path):
    output = tmp_path / "screening.xlsx"
    counts, funds, sectors, excellent, traps, analyzed = generate_excel(
        report_cases["rows"],
        report_cases["hmap"],
        report_cases["evaluations"],
        report_cases["order"],
        output,
        "fixture.xlsx",
        "2026-08-11",
    )

    assert counts == Counter({"Deep Dive": 1, "Watchlist": 1})
    assert (funds, sectors, excellent, traps, analyzed) == (1, 1, 2, 0, 3)
    workbook = openpyxl.load_workbook(output, data_only=False, read_only=False)
    assert workbook.sheetnames == ["Screening", "Instructivo"]
    assert workbook["Screening"]["A3"].value == "🔬 Deep Dive"
    guide_values = [
        workbook["Instructivo"].cell(row=row, column=3).value for row in range(1, 12)
    ]
    assert "Screening cuantitativo — VHC Inversiones" in guide_values
    workbook.close()


def test_excel_flags_missing_optional_columns(evaluate_row, make_case, tmp_path):
    row, header_map, raw, number, percentage = make_case()
    evaluation = evaluate_row(row, raw, number, percentage)
    optional_index = header_map.pop("Beta (5 Year)")
    row.pop(optional_index)
    header_map = {
        header: index if index < optional_index else index - 1
        for header, index in header_map.items()
    }
    output = tmp_path / "opcional.xlsx"

    generate_excel(
        [row],
        header_map,
        {0: evaluation},
        [0],
        output,
        "fixture.xlsx",
        "2026-08-13",
    )

    workbook = openpyxl.load_workbook(output, data_only=False, read_only=False)
    sheet = workbook["Screening"]
    banner = sheet["A1"].value
    assert isinstance(banner, str)
    assert "FALTAN 1 INDICADORES" in banner
    assert "No cambian la clasificación" in banner
    workbook.close()


def test_excel_rejects_missing_core_columns(make_case, tmp_path):
    row, header_map, *_ = make_case()
    header_map.pop("Beneish M-Score")
    output = tmp_path / "invalido.xlsx"

    with pytest.raises(ValueError, match="Beneish M-Score"):
        generate_excel([row], header_map, {}, [], output, "fixture.xlsx", "2026-08-13")

    assert not output.exists()


def test_declining_roic_is_flagged_in_excel(evaluate_row, make_case, tmp_path):
    row, hmap, raw, num, pctv = make_case(
        {
            "Name": "Moat & Co",
            "Ticker": "MOAT",
            "Return on Invested Capital": 0.23,
            "Avg Return on Invested Capital (5y)": 0.24,
            "Fair Value": 130.0,
            "Fair Value Label (Analyst Targets)": "Fair",
            "Free Cash Flow Yield": 0.03,
        }
    )
    result = evaluate_row(row, raw, num, pctv)
    assert result["bucket"] == "Deep Dive"
    assert result["alerts"] == "—"
    output = tmp_path / "moat.xlsx"

    generate_excel(
        [row],
        hmap,
        {0: result},
        [0],
        output,
        "moat.xlsx",
        "2026-08-11",
    )

    workbook = openpyxl.load_workbook(output, data_only=False, read_only=False)
    sheet = workbook["Screening"]
    alert = sheet.cell(row=3, column=_column_for_header(sheet, "Alertas")).value
    assert isinstance(alert, str)
    assert "Revisar ROIC/moat" in alert
    workbook.close()


def test_excel_is_generated_even_without_the_logo(evaluate_row, make_case, tmp_path):
    row, hmap, raw, num, pctv = make_case()
    evaluation = evaluate_row(row, raw, num, pctv)
    output = tmp_path / "sin-logo.xlsx"

    generate_excel(
        [row],
        hmap,
        {0: evaluation},
        [0],
        output,
        "fixture.xlsx",
        "2026-08-11",
        assets_dir=tmp_path / "sin-assets",
    )

    assert output.exists()


def test_excel_neutralizes_formulas_coming_from_the_export(
    evaluate_row, make_case, tmp_path
):
    malicious_name = '=HYPERLINK("https://example.invalid","abrir")'
    row, hmap, raw, num, pctv = make_case({"Name": malicious_name})
    evaluation = evaluate_row(row, raw, num, pctv)
    output = tmp_path / "formula-segura.xlsx"

    generate_excel(
        [row],
        hmap,
        {0: evaluation},
        [0],
        output,
        "fixture.xlsx",
        "2026-08-12",
    )

    workbook = openpyxl.load_workbook(output, data_only=False, read_only=False)
    sheet = workbook["Screening"]
    name_cell = sheet.cell(row=3, column=_column_for_header(sheet, "Name"))
    assert name_cell.value == f"'{malicious_name}"
    assert name_cell.data_type == "s"
    workbook.close()
