#!/usr/bin/env python3
"""Combines several InvestingPro exports into one sheet, verifying continuity.

InvestingPro hands out at most 1,000 rows per export, so a wider universe is
assembled by lowering the market capitalization filter and downloading again.
When that filter is rounded up, the last company of one export reappears as the
first of the next. That duplicate is not a defect: it is the proof that nothing
was skipped between them. A boundary that arrives without one means companies
were lost, and nothing else in the file records their absence.

This tool treats those duplicates as a positive control: it orders the exports
by capitalization, checks every boundary for its overlap, reports the range that
went missing when one is absent, and only then merges and drops the repeats.

Usage:
    python tools/merge_exports.py "~/Downloads/Untitled Screener*.xlsx"
    python tools/merge_exports.py ~/Downloads/Untitled*.xlsx
"""

from __future__ import annotations

import argparse
import glob
import itertools
import math
import os
import sys
from dataclasses import dataclass

import openpyxl
from openpyxl import Workbook

TICKER = "Ticker"
NAME = "Name"
MARKET_CAP = "Market Cap (Adjusted)"
FULL_TICKER = "Full Ticker"

# A genuine InvestingPro export repeats this one column, twice, spelling the
# exchange differently in each — NASDAQGS:NVDA against NasdaqGS:NVDA. Refusing
# duplicates outright would refuse every real file, so this single name is
# tolerated and the last occurrence wins, which is what the classifier reads.
# No other repetition is safe: a second Ticker or Market Cap column would take
# over the first and nothing downstream could tell.
TOLERATED_DUPLICATE = FULL_TICKER
MILLION = 1_000_000

# The ceiling the classifier applies, mirrored here so a merge that would be
# rejected on upload is reported while the files are still on the desk.
# tests/test_merge_exports.py keeps this equal to io.MAX_WORKSHEET_ROWS.
CLASSIFIER_ROW_LIMIT = 10_000


def _text(value: object) -> str:
    """Collapses the whitespace and line breaks a cell may carry."""
    return " ".join(str(value).split()) if value is not None else ""


def _number(value: object) -> float | None:
    """Returns the value as a float only when it is genuinely numeric."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


@dataclass
class Batch:
    """One downloaded export, read and indexed by column name."""

    path: str
    header: list[str]
    columns: dict[str, int]
    rows: list[list[object]]
    tickers: list[str]
    identities: list[str]
    caps: list[float]

    @property
    def name(self) -> str:
        return os.path.basename(self.path)

    @property
    def top(self) -> float:
        return max(self.caps, default=0.0)

    @property
    def bottom(self) -> float:
        return min(self.caps, default=0.0)


@dataclass(frozen=True)
class Gap:
    """A boundary that arrived with no overlap, and the range it lost."""

    after: str
    before: str
    lost_from: float
    lost_to: float

    @property
    def resume_at(self) -> float:
        """The filter that recovers the range, in millions, rounded up.

        Rounding up is what guarantees the overlap; rounding down is what opens
        the gap, because the boundary company then falls outside both exports.
        """
        return math.ceil(self.lost_to / MILLION * 100) / 100


def _read_sheet(path: str) -> list[list[object]]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book.active
    if sheet is None:
        raise SystemExit(f"ERROR: {os.path.basename(path)} no trae ninguna hoja")
    rows: list[list[object]] = [list(row) for row in sheet.iter_rows(values_only=True)]
    book.close()
    return rows


def _find_header(rows: list[list[object]], path: str) -> int:
    """Locates the header row wherever it sits, by its labels."""
    for index, row in enumerate(rows):
        labels = {_text(cell) for cell in row}
        if NAME in labels and TICKER in labels:
            return index
    raise SystemExit(f"ERROR: {os.path.basename(path)} no trae fila de encabezados")


def _identity(row: list[object], columns: dict[str, int]) -> str:
    """What makes a company itself, which is not its ticker.

    Two companies really do share one: in an August 2026 export, `FLYX` is both
    Flyexclusive Inc on NYSE and FlyE Group Inc on NYSEAM. Deduplicating by
    ticker drops the second without a word, and the merged sheet gives no sign
    that a company went missing.

    `Full Ticker` carries the exchange and settles it. It is optional for the
    method, so when it is absent the pair (ticker, name) stands in — enough to
    tell two companies apart while still recognising the same one across two
    exports.
    """
    full = _text(_cell(row, columns.get(FULL_TICKER))).casefold()
    if full:
        return full
    ticker = _text(_cell(row, columns.get(TICKER))).casefold()
    name = _text(_cell(row, columns.get(NAME))).casefold()
    return f"{ticker}\x00{name}"


def _cell(row: list[object], index: int | None) -> object:
    if index is None or index >= len(row):
        return None
    return row[index]


def _require_boundary_capitalizations(
    data: list[list[object]], name_at: int, cap_at: int, path: str
) -> None:
    """The first and last companies decide the order and the next filter.

    A gap in the middle is harmless: capitalization is optional for the method,
    and those rows merge like any other. The edges are not. The batch is placed
    among the others by its first row and the next filter is computed from its
    last, so a missing value there pairs one company's ticker with another
    company's number — and the tool would announce a boundary that never existed.
    """
    if not data:
        raise SystemExit(f"ERROR: {os.path.basename(path)} no trae ninguna empresa")
    for row, edge in ((data[0], "primera"), (data[-1], "última")):
        if _number(_cell(row, cap_at)) is not None:
            continue
        company = _text(_cell(row, name_at))
        raise SystemExit(
            f"ERROR: la {edge} empresa de {os.path.basename(path)} — {company} — "
            f"no trae '{MARKET_CAP}'. Ese valor sitúa la tanda entre las demás y "
            "calcula el filtro siguiente, así que sin él la frontera que se "
            "anuncie sería de otra empresa."
        )


def _require_descending_order(
    data: list[list[object]], ticker_at: int, cap_at: int, path: str
) -> None:
    """The export must run from the largest capitalization down.

    Everything the tool reports about a batch reads its edges: the first row
    places it among the others, the last one yields the next filter. Both are
    read as tickers while the figures come from the whole column, so a sheet
    sorted any other way pairs one company's name with another's number and
    announces a boundary that never existed.

    Refused rather than re-sorted. Sorting would hide the real cause — a screener
    exported without `Market Cap (Adjusted)` descending — and the next download
    would repeat it.
    """
    previous_value: float | None = None
    previous_ticker = ""
    for index, row in enumerate(data, 1):
        value = _number(_cell(row, cap_at))
        if value is None:  # An interior gap, already reported elsewhere.
            continue
        if previous_value is not None and value > previous_value:
            ticker = _text(_cell(row, ticker_at))
            raise SystemExit(
                f"ERROR: {os.path.basename(path)} no está ordenado de mayor a "
                f"menor capitalización: {ticker} ({value / MILLION:,.2f} M) en "
                f"la fila {index} supera a {previous_ticker} "
                f"({previous_value / MILLION:,.2f} M), que va antes. Vuelve a "
                f"exportar ordenando por '{MARKET_CAP}' de forma descendente."
            )
        previous_value, previous_ticker = value, _text(_cell(row, ticker_at))


def _check_repeated_headers(header_row: list[object], path: str) -> None:
    """Only `Full Ticker` may appear twice; anything else is a silent swap."""
    seen: dict[str, list[int]] = {}
    for index, cell in enumerate(header_row, 1):
        label = _text(cell)
        if label:
            seen.setdefault(label, []).append(index)

    name = os.path.basename(path)
    for label, positions in sorted(seen.items()):
        if len(positions) == 1:
            continue
        if label == TOLERATED_DUPLICATE and len(positions) == 2:
            print(
                f"  aviso: {name} repite '{label}' en las columnas "
                f"{positions[0]} y {positions[1]}; se usa la última, igual que "
                "el clasificador"
            )
            continue
        if label == TOLERATED_DUPLICATE:
            # Two is what InvestingPro writes. More than that is a file nobody
            # produced on purpose, and the last column decides identity, so a
            # third would quietly change which companies count as the same.
            raise SystemExit(
                f"ERROR: {name} trae '{label}' {len(positions)} veces, en las "
                f"posiciones {', '.join(map(str, positions))}. Un export "
                "auténtico la trae dos, y la última decide la identidad de cada "
                "empresa: con más, esa elección deja de ser predecible."
            )
        raise SystemExit(
            f"ERROR: {name} repite la columna '{label}' en las posiciones "
            f"{', '.join(map(str, positions))}. Solo '{TOLERATED_DUPLICATE}' "
            "puede venir duplicada; cualquier otra tapa a la primera en "
            "silencio y el dato leído sería el equivocado."
        )


def read_batch(path: str) -> Batch:
    """Reads one export and pulls out its tickers and capitalizations."""
    rows = _read_sheet(path)
    header_at = _find_header(rows, path)
    _check_repeated_headers(rows[header_at], path)
    columns = {_text(v): j for j, v in enumerate(rows[header_at]) if _text(v)}
    name_at, ticker_at = columns[NAME], columns[TICKER]
    if MARKET_CAP not in columns:
        raise SystemExit(
            f"ERROR: {os.path.basename(path)} no trae la columna "
            f"'{MARKET_CAP}'. Sin ella no se puede ordenar las tandas ni "
            "comprobar que no falten empresas entre una y otra."
        )
    cap_at = columns[MARKET_CAP]

    data: list[list[object]] = []
    tickers: list[str] = []
    identities: list[str] = []
    caps: list[float] = []
    for row in rows[header_at + 1 :]:
        if not row or not _text(_cell(row, name_at)):
            continue
        data.append(list(row))
        tickers.append(_text(_cell(row, ticker_at)))
        identities.append(_identity(row, columns))
        capitalization = _number(_cell(row, cap_at))
        if capitalization is not None:
            caps.append(capitalization)

    header = [_text(cell) for cell in rows[header_at]]
    _require_boundary_capitalizations(data, name_at, cap_at, path)
    _require_descending_order(data, ticker_at, cap_at, path)
    if data and len(caps) < len(data):
        # Harmless for the classification — the method treats capitalization as
        # optional — but these rows sit outside the top and bottom that order
        # the batches, so the reader should know they are there.
        print(
            f"  aviso: {os.path.basename(path)} trae {len(data) - len(caps)} "
            f"filas sin '{MARKET_CAP}'; no cuentan para ordenar las tandas"
        )
    return Batch(
        path=path,
        header=header,
        columns=columns,
        rows=data,
        tickers=tickers,
        identities=identities,
        caps=caps,
    )


def find_gaps(batches: list[Batch]) -> list[Gap]:
    """Every boundary must repeat a company. Its absence is a gap."""
    gaps: list[Gap] = []
    for current, following in itertools.pairwise(batches):
        if set(current.identities) & set(following.identities):
            continue
        gaps.append(
            Gap(
                after=current.name,
                before=following.name,
                lost_from=following.top,
                lost_to=current.bottom,
            )
        )
    return gaps


def require_matching_columns(batches: list[Batch]) -> None:
    """Every export must carry the same columns, or the merge is meaningless.

    A screener reconfigured between downloads yields a different column set, and
    a merged sheet built from both would hold blanks where the method expects
    figures. Stopping names the file and what differs; carrying on would classify
    those companies on data that is not there.
    """
    canonical = set(batches[0].columns)
    for batch in batches[1:]:
        missing = sorted(canonical - set(batch.columns))
        extra = sorted(set(batch.columns) - canonical)
        if not missing and not extra:
            continue
        detail = []
        if missing:
            detail.append(f"le faltan: {', '.join(missing)}")
        if extra:
            detail.append(f"trae de más: {', '.join(extra)}")
        raise SystemExit(
            f"ERROR: {batch.name} no tiene las mismas columnas que "
            f"{batches[0].name} — {'; '.join(detail)}. Vuelve a exportar todas "
            "las tandas con la misma configuración del screener."
        )


def merge(batches: list[Batch]) -> tuple[list[str], list[list[object]], list[str]]:
    """Merges the rows, keeping the first appearance of every company.

    Kept by identity, not by ticker: two companies can share one, and keying on
    it drops the second without a word. See `_identity`.

    Values are read by column name, never by position: two exports may list the
    same columns in a different order, and lining them up by position would slide
    every value one cell sideways without saying so.
    """
    require_matching_columns(batches)
    canonical = [
        name for name, _ in sorted(batches[0].columns.items(), key=lambda kv: kv[1])
    ]
    seen: set[str] = set()
    merged: list[list[object]] = []
    repeated: list[str] = []
    for batch in batches:
        for ticker, identity, row in zip(
            batch.tickers, batch.identities, batch.rows, strict=True
        ):
            if identity in seen:
                repeated.append(ticker)
                continue
            seen.add(identity)
            merged.append([_cell(row, batch.columns.get(name)) for name in canonical])
    return canonical, merged, repeated


def write_sheet(header: list[str], rows: list[list[object]], output: str) -> None:
    """Writes the merged universe under a single header row."""
    book = Workbook()
    sheet = book.active
    if sheet is None:  # pragma: no cover - a fresh workbook always has a sheet
        raise SystemExit("ERROR: no se pudo crear la hoja de salida")
    sheet.title = "Universo"
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    book.save(output)


def _print_batches(batches: list[Batch]) -> None:
    print("== TANDAS, ordenadas por capitalización ==")
    for batch in batches:
        first = batch.tickers[0] if batch.tickers else ""
        last = batch.tickers[-1] if batch.tickers else ""
        print(
            f"  {len(batch.rows):>5,} filas · "
            f"{batch.top / MILLION:>14,.2f} M -> {batch.bottom / MILLION:>12,.2f} M · "
            f"{first:>6} .. {last:<6} · {batch.name}"
        )


def _print_continuity(batches: list[Batch], gaps: list[Gap]) -> None:
    # A single export has no boundary to check, so there is nothing to claim
    # about it either way.
    if len(batches) < 2:
        return
    print("\n== CONTINUIDAD ENTRE TANDAS ==")
    lost = {gap.after for gap in gaps}
    for current, following in itertools.pairwise(batches):
        if current.name in lost:
            print(
                f"  HUECO entre {following.top / MILLION:,.2f} M y "
                f"{current.bottom / MILLION:,.2f} M  <-- faltan empresas"
            )
            continue
        overlap = set(current.identities) & set(following.identities)
        shared = sorted(
            {
                ticker
                for ticker, identity in zip(
                    current.tickers, current.identities, strict=True
                )
                if identity in overlap
            }
        )
        print(
            f"  OK   {current.bottom / MILLION:>12,.2f} M  ->  "
            f"solape: {', '.join(shared[:3])}"
        )


def _print_next_step(batches: list[Batch], rows_per_export: int = 1000) -> None:
    """Says what to type next, so the smallest batch need not be read by eye."""
    smallest = batches[-1]
    if len(smallest.rows) < rows_per_export:
        print(
            f"\nÚltima tanda con {len(smallest.rows):,} filas "
            f"(menos de {rows_per_export:,}): el universo está completo."
        )
        return
    following = math.ceil(smallest.bottom / MILLION * 100) / 100
    print(
        f"\n== SIGUIENTE TANDA ==\n"
        f"  La última descargada acaba en {smallest.tickers[-1]}, "
        f"{smallest.bottom / MILLION:,.2f} M. Para la siguiente:\n"
        f"      Market Cap (Adjusted) · less than or equal to · "
        f"{following} · million\n"
        f"  Si {following} ya está por debajo del tamaño que te interesa, "
        f"has terminado: deja de descargar."
    )


def _require_a_sheet_the_classifier_accepts(companies: int) -> None:
    """The ceiling counts the header, so one company fewer than it fits.

    Checked before writing: a sheet the screening will refuse is not worth
    putting on disk, and hearing it here beats hearing it after the upload.
    """
    rows_in_sheet = 1 + companies
    if rows_in_sheet <= CLASSIFIER_ROW_LIMIT:
        return
    raise SystemExit(
        f"ERROR: {companies:,} empresas más la fila de encabezados hacen "
        f"{rows_in_sheet:,} filas, y el clasificador acepta "
        f"{CLASSIFIER_ROW_LIMIT:,}. Sube el piso de capitalización y vuelve a "
        "fusionar: el archivo no se ha escrito."
    )


def _print_summary(
    batches: list[Batch], merged: list[list[object]], repeated: list[str], output: str
) -> None:
    unique = list(dict.fromkeys(repeated))
    print("\n== FUSIÓN ==")
    print(f"  filas leídas    : {sum(len(b.rows) for b in batches):,}")
    print(f"  duplicados      : {len(repeated)}  {unique[:12] if unique else ''}")
    print(f"  empresas únicas : {len(merged):,}")
    print(f"  archivo         : {output}")


def _print_gaps(gaps: list[Gap], boundaries: int) -> None:
    if boundaries == 0:
        return
    if not gaps:
        print("\nSin huecos: cada frontera trae su solape.")
        return
    print("\n== HUECOS QUE HAY QUE CERRAR ==")
    for gap in gaps:
        print(
            f"  - {gap.after} -> {gap.before}: sin solape. "
            f"Se perdió todo entre {gap.lost_from / MILLION:,.2f} M "
            f"y {gap.lost_to / MILLION:,.2f} M. Vuelve a descargar con:"
        )
        # The web filter lets the user pick the unit; the API only speaks
        # millions. Both are printed so neither workflow has to convert.
        print(
            f"      en la web  ->  Market Cap (Adjusted) · menor o igual que · "
            f"{gap.resume_at} · million"
        )
        print(f'      en el curl ->  "marketcap_adj": {{"$lte": {gap.resume_at}}}')


def resolve_paths(patterns: list[str]) -> list[str]:
    """Resolves the arguments, whether the shell expanded the wildcard or not.

    Quoting the pattern hands this one string with a `*` in it; leaving it
    unquoted hands one argument per file. Both are accepted so the command
    works either way, and duplicates are dropped in case the two forms overlap.
    """
    found: list[str] = []
    for pattern in patterns:
        expanded = os.path.expanduser(pattern)
        found.extend(glob.glob(expanded) if glob.has_magic(expanded) else [expanded])
    unique = dict.fromkeys(p for p in found if not os.path.basename(p).startswith("~$"))
    return sorted(unique)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "patterns",
        nargs="+",
        metavar="EXPORT",
        help="Rutas o comodín de los exports a fusionar",
    )
    parser.add_argument(
        "-o",
        "--output",
        default=None,
        help="Archivo de salida. Por omisión, universe.xlsx junto a los exports",
    )
    args = parser.parse_args(argv)

    paths = resolve_paths(args.patterns)
    if not paths:
        raise SystemExit(
            f"ERROR: ningún archivo coincide con {' '.join(args.patterns)}"
        )

    # Largest capitalization first: the real download order, not the file names.
    batches = sorted((read_batch(p) for p in paths), key=lambda b: b.top, reverse=True)

    _print_batches(batches)
    gaps = find_gaps(batches)
    _print_continuity(batches, gaps)

    header, merged, repeated = merge(batches)
    _require_a_sheet_the_classifier_accepts(len(merged))
    # Written next to the exports it merged, not into whatever directory the
    # command happened to run from: the result belongs with its sources.
    output = (
        os.path.expanduser(args.output)
        if args.output
        else os.path.join(os.path.dirname(paths[0]), "universe.xlsx")
    )
    write_sheet(header, merged, output)

    _print_summary(batches, merged, repeated, output)
    _print_gaps(gaps, len(batches) - 1)
    _print_next_step(batches)
    return 1 if gaps else 0


if __name__ == "__main__":
    sys.exit(main())
